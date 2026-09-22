from flask import Flask, render_template, request, jsonify, send_file, send_from_directory, Response, stream_with_context, redirect, url_for
from utils import rag_pipeline, process_triple, wrap_text
from case_generator import CaseClassifier, ReverseTestCaseGenerator
from llm_client import generate_testcase
from skills import skill_registry
from storage.review_store import review_store
import json
import os
import pandas as pd
from datetime import datetime, timezone
import io
from flask_cors import CORS
import re
import queue
import threading
import time
import uuid
from pathlib import Path
from werkzeug.utils import secure_filename

app = Flask(__name__)

app.config['JSON_AS_ASCII'] = False
CORS(app)

# Initialize case classifier and generator
case_classifier = CaseClassifier('data/TestCase_v2.json')
reverse_generator = ReverseTestCaseGenerator()
UPLOAD_DIR = Path("data/uploads/sts")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

def get_kg_workbench_url():
    return os.getenv("KG_WORKBENCH_URL", "http://127.0.0.1:8501")

@app.route('/')
def index():
    """Home page: prefer root index.html so Live Preview / GitHub Pages / Flask share one entry."""
    root_index = Path('index.html')
    if root_index.exists():
        return send_from_directory(Path.cwd(), 'index.html')
    return render_template('platform.html', kg_workbench_url=get_kg_workbench_url())


@app.route('/index.html')
def index_html():
    """Allow opening /index.html directly for the same full workbench entry."""
    root_index = Path('index.html')
    if root_index.exists():
        return send_from_directory(Path.cwd(), 'index.html')
    return redirect('/')


@app.route('/qa')
def qa_page():
    """Keep QA inside the main workbench shell (with left rail)."""
    return redirect('/#qa')


@app.route('/test-case-generation')
def test_case_generation_page():
    """Deep-link into the native test-case generation workbench panel."""
    return redirect('/#test-case-generation')


@app.route('/reverse-generation')
def reverse_generation():
    """Batch reverse test-case generation workbench."""
    return render_template('reverse_generation.html')


@app.route('/fault-diagnosis')
def fault_diagnosis():
    """Fault cause analysis workflow page."""
    return render_template('fault_diagnosis.html')


@app.route('/api/fault-diagnosis/stream', methods=['POST'])
def fault_diagnosis_stream():
    """Stream a lightweight local fault-diagnosis analysis over related cases."""
    payload = request.get_json(silent=True) or {}
    problem = str(payload.get('problem') or '').strip()
    if not problem:
        return jsonify({"success": False, "error": "Please describe the fault symptom"}), 400

    def generate():
        started = time.time()
        yield json.dumps({"type": "thought", "message": "Retrieving test case evidence for the current product line..."}, ensure_ascii=False) + "\n"
        try:
            cases = _qa_retrieve_cases(problem, limit=6)
        except RuntimeError as exc:
            yield json.dumps({"type": "error", "message": str(exc)}, ensure_ascii=False) + "\n"
            return

        evidence = []
        for index, item in enumerate(cases, start=1):
            evidence_id = f"E{index}"
            statement_bits = [
                item.get('case_name') or '',
                '；'.join(item.get('preconditions') or [])[:120],
                '；'.join(item.get('actions') or [])[:120],
                '；'.join(item.get('expected_behaviors') or [])[:120],
            ]
            evidence.append({
                "id": evidence_id,
                "source_type": "testcase",
                "source_name": item.get('source') or "TestCase_v2.json",
                "source_row": item.get('source_row'),
                "score": item.get('score'),
                "statement": " | ".join([bit for bit in statement_bits if bit]) or item.get('case_name') or evidence_id,
                "case_id": item.get('case_id'),
                "case_name": item.get('case_name'),
                "test_type": item.get('test_type'),
            })

        yield json.dumps({"type": "thought", "message": f"Retrieved {len(evidence)} related case evidence item(s); inferring possible causes..."}, ensure_ascii=False) + "\n"
        yield json.dumps({"type": "evidence", "data": evidence}, ensure_ascii=False) + "\n"

        causes = []
        for index, item in enumerate(cases[:4], start=1):
            likelihood = 'high' if index == 1 else ('medium' if index == 2 else 'low')
            checks = []
            if item.get('preconditions'):
                checks.append(f"Verify preconditions: {'; '.join(item['preconditions'][:2])}")
            if item.get('actions'):
                checks.append(f"Reproduce actions: {'; '.join(item['actions'][:2])}")
            if item.get('expected_behaviors'):
                checks.append(f"Compare expected behavior: {'; '.join(item['expected_behaviors'][:2])}")
            causes.append({
                "title": item.get('case_name') or f"Related scenario {index}",
                "likelihood": likelihood,
                "reasoning": f"Test type related to the fault: \"{item.get('test_type') or '未分类'}\"; use as an investigation entry point.",
                "checks": checks or ["Cross-check on-site signals and logs."],
                "evidence_ids": [f"E{index}"],
            })

        if not causes:
            causes = [{
                "title": "Insufficient evidence; add fault context",
                "likelihood": "low",
                "reasoning": "Not enough matching test cases; add system, conditions, and symptom details.",
                "checks": ["Add fault conditions", "Add controller/function names", "Note cluster warnings"],
                "evidence_ids": [],
            }]

        next_steps = [
            "Run quick checks for highest-likelihood causes",
            "Reproduce fault using preconditions/actions from evidence",
            "If unresolved, add conditions and re-analyze",
        ]
        result = {
            "summary": f"Initial fault-cause analysis for \"{problem}\" based on test cases.",
            "causes": causes,
            "evidence": evidence,
            "next_steps": next_steps,
            "insufficient_evidence": len(evidence) < 2,
            "elapsed_time": round(time.time() - started, 2),
        }
        yield json.dumps({"type": "thought", "message": "Analysis complete"}, ensure_ascii=False) + "\n"
        yield json.dumps({"type": "result", "data": result}, ensure_ascii=False) + "\n"

    return Response(stream_with_context(generate()), mimetype='application/x-ndjson')


@app.route('/legacy-generation')
def legacy_generation():
    """Legacy RAG test-case generation workbench (workflows entry)."""
    return render_template('legacy_generation.html', kg_workbench_url=get_kg_workbench_url())


@app.route('/type-expected-classification')
def type_expected_classification():
    """Compatibility alias for the reverse-generation workbench."""
    return redirect(url_for('reverse_generation'))


@app.route('/api/auth/session', methods=['GET'])
def auth_session():
    """Local preview session — ontology workbench requires authenticated ready state."""
    return jsonify({"success": True, "data": {
        "authenticated": True,
        "user": {"username": "local", "display_name": "Local user", "role": "admin", "status": "active"},
        "permissions": ["kg:read", "kg:write"],
        "knowledge_graphs": [{"id": "testcase", "name": "Test case knowledge graph"}],
        "current_kg_id": "testcase",
    }})


@app.route('/api/auth/login', methods=['POST'])
def auth_login_placeholder():
    """Stable login contract until the account service is connected."""
    return jsonify({"success": False, "error": "Account login service is not connected yet", "code": "AUTH_NOT_CONFIGURED"}), 501


@app.route('/api/knowledge-graphs', methods=['GET'])
def list_knowledge_graphs():
    """Return the currently available KG until account-scoped filtering is connected."""
    return jsonify({"success": True, "data": [{
        "id": "testcase",
        "name": "Test case knowledge graph",
        "domain": "Vehicle test cases",
        "role": "viewer",
    }], "current_kg_id": "testcase", "source": "current-dataset"})


@app.route('/api/graph/data', methods=['GET'])
def graph_data():
    """Expose a compact graph projection aligned with the remote #graph contract."""
    kg_id = request.args.get('kg_id', 'testcase')
    case_id = (request.args.get('case_id') or '').strip()
    limit_raw = request.args.get('limit', '100')
    if str(limit_raw).lower() == 'all':
        limit = 240
    else:
        try:
            limit = int(limit_raw)
        except (TypeError, ValueError):
            limit = 100
    limit = min(max(limit, 20), 240)

    try:
        cases = _load_search_cases()
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 500

    total_cases = len(cases)
    total_types = len({item.get('test_type') for item in cases if item.get('test_type')})

    if case_id:
        selected = [item for item in cases if item.get('case_id') == case_id or item.get('id') == case_id]
        if not selected:
            selected = [item for item in cases if case_id in (item.get('case_name') or '')]
        sample = selected[:1]
    else:
        # Round-robin across test types so overview is a mesh, not one type hub.
        by_type = {}
        for item in cases:
            test_type = item.get('test_type') or '未分类'
            by_type.setdefault(test_type, []).append(item)
        type_names = sorted(by_type.keys(), key=lambda name: (-len(by_type[name]), name))
        sample = []
        cursor = {name: 0 for name in type_names}
        while len(sample) < limit and type_names:
            progressed = False
            for name in list(type_names):
                index = cursor[name]
                bucket = by_type[name]
                if index >= len(bucket):
                    type_names.remove(name)
                    continue
                sample.append(bucket[index])
                cursor[name] = index + 1
                progressed = True
                if len(sample) >= limit:
                    break
            if not progressed:
                break

    nodes = {}
    edges = []
    triple_index = 0

    def add_node(node_id, label, kind, **extra):
        if node_id in nodes:
            nodes[node_id]['degree'] = nodes[node_id].get('degree', 0) + 1
            return nodes[node_id]
        node = {"id": node_id, "label": label, "kind": kind, "degree": 1}
        node.update(extra)
        nodes[node_id] = node
        return node

    def add_edge(source, target, relation, triple_type='测试用例业务关系'):
        nonlocal triple_index
        edge_id = f"triple:{triple_index}"
        triple_index += 1
        edges.append({
            "id": edge_id,
            "source": source,
            "target": target,
            "label": relation,
            "relation": relation,
            "triple_type": triple_type,
        })

    # Remote overview at limit=60 reports 351 nodes / 291 relations / 60 types / 60 cases
    # → 231 detail nodes with one edge each. Keep the same density for other limits.
    detail_budget = None if case_id else max(0, int(round(len(sample) * 3.85)))
    detail_count = 0

    for item in sample:
        test_type = item.get('test_type') or '未分类'
        type_id = f"type:{test_type}"
        case_key = item.get('case_id') or item.get('id')
        case_node_id = f"case:{case_key}"
        add_node(type_id, test_type, 'type', size=22)
        add_node(
            case_node_id,
            item.get('case_name') or case_key,
            'case',
            case_type=test_type,
            case_number=case_key,
            size=14,
        )
        add_edge(type_id, case_node_id, '包含用例')

        detail_slots = []
        for index, precondition in enumerate(item.get('preconditions') or []):
            detail_slots.append(('condition', f"condition:{case_key}:{index}", precondition, '前提条件'))
        for index, action in enumerate(item.get('actions') or []):
            detail_slots.append(('action', f"action:{case_key}:{index}", action, '执行动作'))
        for index, expected in enumerate(item.get('expected_behaviors') or []):
            detail_slots.append(('expectation', f"expectation:{case_key}:{index}", expected, '预期行为'))

        # Remote #graph (limit=60) shows 351 nodes / 291 edges = 231 detail
        # triples. Match that density: ≤4 details per case, global budget ≈ 3.85×cases.
        if case_id:
            selected_details = detail_slots
        else:
            selected_details = detail_slots[:4]
            room = detail_budget - detail_count
            if room <= 0:
                selected_details = []
            elif len(selected_details) > room:
                selected_details = selected_details[:room]

        for kind, node_id, label, relation in selected_details:
            add_node(node_id, label, kind, size=9)
            add_edge(case_node_id, node_id, relation)
            detail_count += 1

    ordered_nodes = list(nodes.values())
    type_count = sum(1 for node in ordered_nodes if node['kind'] == 'type')
    case_count = sum(1 for node in ordered_nodes if node['kind'] == 'case')
    return jsonify({
        "success": True,
        "data": {
            "kg_id": kg_id,
            "nodes": ordered_nodes,
            "edges": edges,
            "stats": {
                "nodes": len(ordered_nodes),
                "edges": len(edges),
                "types": type_count,
                "cases": case_count,
                "categories": type_count,
                "triples": len(edges),
                "total_cases": total_cases,
                "total_types": total_types,
                "source": "testcase-dataset",
            },
        },
    })


@app.route('/api/graph/list', methods=['GET'])
def graph_list():
    """List projection for the graph list mode."""
    kg_id = request.args.get('kg_id', 'testcase')
    type_filter = (request.args.get('type') or '').strip()
    query = (request.args.get('q') or '').strip().lower()
    limit_raw = request.args.get('limit', '180')
    try:
        cases = _load_search_cases()
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 500

    def case_id_key(value):
        text = str(value or '')
        match = re.search(r'(\d+)$', text)
        return (int(match.group(1)) if match else 10 ** 12, text)

    filtered = []
    for item in cases:
        test_type = item.get('test_type') or '未分类'
        if type_filter and test_type != type_filter:
            continue
        blob = ' '.join([
            item.get('case_id') or '',
            item.get('case_name') or '',
            test_type,
            ' '.join(item.get('preconditions') or []),
            ' '.join(item.get('actions') or []),
            ' '.join(item.get('expected_behaviors') or []),
        ]).lower()
        if query and query not in blob:
            continue
        filtered.append({
            "id": item.get('case_id') or item.get('id'),
            "case_number": item.get('case_id') or item.get('id'),
            "label": item.get('case_name') or item.get('case_id'),
            "type": test_type,
            "condition_count": len(item.get('preconditions') or []),
            "has_action": bool(item.get('actions')),
            "has_expectation": bool(item.get('expected_behaviors')),
        })

    by_type = {}
    for record in filtered:
        by_type.setdefault(record['type'], []).append(record)
    for records in by_type.values():
        records.sort(key=lambda row: case_id_key(row.get('id')))
    type_order = sorted(by_type.keys(), key=lambda name: case_id_key(by_type[name][0].get('id')))
    available_types = [{"name": name, "count": len(by_type[name])} for name in type_order]

    # Stats follow the remote contract: summarize the full filter match,
    # then select display rows by type min-id order + round-robin.
    matched_count = len(filtered)
    type_count = len(type_order)
    stats = {
        "types": type_count,
        "total_cases": matched_count,
        "relations": matched_count,
        "average_cases_per_type": round(matched_count / type_count, 1) if type_count else 0,
    }

    if str(limit_raw).lower() == 'all':
        selected = [record for name in type_order for record in by_type[name]]
    else:
        try:
            limit = max(1, int(limit_raw))
        except (TypeError, ValueError):
            limit = 180
        queues = {name: list(by_type[name]) for name in type_order}
        selected = []
        while len(selected) < limit and any(queues[name] for name in type_order):
            for name in type_order:
                if len(selected) >= limit:
                    break
                if queues[name]:
                    selected.append(queues[name].pop(0))

    stats["cases"] = len(selected)

    return jsonify({
        "success": True,
        "data": {
            "kg_id": kg_id,
            "cases": selected,
            "types": available_types,
            "available_types": available_types,
            "stats": stats,
        },
    })


_SEARCH_CASES_CACHE = {"mtime": None, "cases": []}


def _as_text_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    parts = [part.strip() for part in re.split(r'[\n；;]+', text) if part.strip()]
    return parts or [text]


def _load_search_cases():
    source = Path('data/TestCase_v2.json')
    try:
        mtime = source.stat().st_mtime
    except OSError as exc:
        raise RuntimeError(f"Search data unavailable: {exc}") from exc
    if _SEARCH_CASES_CACHE["mtime"] == mtime and _SEARCH_CASES_CACHE["cases"]:
        return _SEARCH_CASES_CACHE["cases"]
    try:
        data = json.loads(source.read_text(encoding='utf-8'))
        cases = data.get('test_cases', data if isinstance(data, list) else [])
    except (OSError, ValueError) as exc:
        raise RuntimeError(f"Search data unavailable: {exc}") from exc

    normalized = []
    for index, item in enumerate(cases):
        if not isinstance(item, dict):
            continue
        original = item.get('original_data', item)
        if not isinstance(original, dict):
            original = {}
        case_id = str(original.get('编号') or f'case-{index}')
        case_name = str(original.get('测试用例名称') or case_id)
        test_type = str(original.get('测试用例类型') or '')
        actions = _as_text_list(original.get('执行动作'))
        expected = _as_text_list(original.get('预期行为'))
        preconditions = _as_text_list(item.get('前提条件_分割') or original.get('前提条件'))
        blob = ' '.join([
            case_id,
            case_name,
            test_type,
            ' '.join(actions),
            ' '.join(expected),
            ' '.join(preconditions),
        ]).lower()
        normalized.append({
            "id": case_id,
            "case_id": case_id,
            "case_number": case_id,
            "label": case_name,
            "case_name": case_name,
            "test_type": test_type,
            "case_type": test_type,
            "preconditions": preconditions,
            "actions": actions,
            "expected_behaviors": expected,
            "source": "TestCase_v2.json",
            "source_row": index + 1,
            "_blob": blob,
        })
    _SEARCH_CASES_CACHE["mtime"] = mtime
    _SEARCH_CASES_CACHE["cases"] = normalized
    return normalized


@app.route('/api/graph/search', methods=['GET', 'POST'])
def graph_search():
    """Search test cases by type, name, action, or expected behavior."""
    payload = request.get_json(silent=True) or {}
    kg_id = request.args.get('kg_id') or payload.get('kg_id') or 'testcase'
    query = (request.args.get('q') or payload.get('query') or payload.get('q') or '').strip()
    limit = min(max(int(request.args.get('limit') or payload.get('limit') or 50), 1), 100)

    if not query:
        return jsonify({"success": True, "data": {"kg_id": kg_id, "query": query, "results": [], "total": 0}, "source": "testcase-dataset"})

    try:
        cases = _load_search_cases()
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 500

    needle = query.lower()
    results = []
    for item in cases:
        if needle not in item["_blob"]:
            continue
        case_name = item["case_name"]
        test_type = item["test_type"]
        case_id = item["case_id"]
        score = 0.0
        if needle in case_name.lower():
            score += 0.55
        if needle in test_type.lower():
            score += 0.25
        if needle in case_id.lower():
            score += 0.15
        if needle in ' '.join(item["actions"]).lower() or needle in ' '.join(item["expected_behaviors"]).lower():
            score += 0.1
        if needle in ' '.join(item["preconditions"]).lower():
            score += 0.05
        row = {key: value for key, value in item.items() if key != "_blob"}
        row["score"] = round(min(score, 0.999), 4)
        results.append(row)

    results.sort(key=lambda row: (-row.get('score', 0), row.get('case_name') or ''))
    total = len(results)
    results = results[:limit]
    return jsonify({
        "success": True,
        "data": {"kg_id": kg_id, "query": query, "results": results, "total": total},
        "source": "testcase-dataset",
    })


@app.route('/api/ontology', methods=['GET'])
def ontology_placeholder():
    """Standard ontology contract reserved for permission-scoped KG data."""
    return jsonify({"success": True, "data": {"kg_id": request.args.get('kg_id'), "classes": [], "properties": []}, "source": "placeholder"})


@app.route('/api/kg/sources', methods=['GET'])
def kg_sources():
    """Serve ontology source registry aligned with remote #ontology."""
    snapshot = Path('data/kg_sources.json')
    if snapshot.exists():
        try:
            payload = json.loads(snapshot.read_text(encoding='utf-8'))
            if isinstance(payload, dict) and payload.get('data'):
                data = payload['data']
                if isinstance(data, dict) and 'trash_count' not in data:
                    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
                    data = dict(data)
                    data['trash_count'] = len(trash)
                return jsonify({"success": True, "data": data})
            if isinstance(payload, dict) and ('excel' in payload or 'sources' in payload):
                data = dict(payload)
                trash = data.get('trash') if isinstance(data.get('trash'), list) else []
                data['trash_count'] = len(trash)
                return jsonify({"success": True, "data": data})
        except (OSError, ValueError):
            pass

    cases = []
    try:
        cases = _load_search_cases()
    except RuntimeError:
        cases = []
    type_count = len({item.get('test_type') for item in cases if item.get('test_type')})
    excel = [{
        "id": "testcase-v2",
        "source_id": "testcase-v2",
        "kind": "excel",
        "name": "TestCase_v2.json",
        "domain": "Test cases",
        "status": "available",
        "case_count": len(cases),
        "triple_count": len(cases) * 3,
        "triples": len(cases) * 3,
        "version": 1,
        "bytes": Path('data/TestCase_v2.json').stat().st_size if Path('data/TestCase_v2.json').exists() else 0,
    }]
    return jsonify({
        "success": True,
        "data": {
            "excel": excel,
            "pdf": [],
            "sources": excel,
            "summary": {"case_count": len(cases), "type_count": type_count},
            "graph": {"case_count": len(cases), "file": "TestCase_v2.json", "total": len(cases), "triple_count": len(cases) * 3, "version": 0},
        },
        "source": "testcase-dataset",
    })


def _kg_load_sources_payload():
    snapshot = Path('data/kg_sources.json')
    if not snapshot.exists():
        return None, None
    try:
        payload = json.loads(snapshot.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return None, None
    data = payload.get('data') if isinstance(payload, dict) and isinstance(payload.get('data'), dict) else payload
    if not isinstance(data, dict):
        return None, None
    return payload, data


def _kg_save_sources_payload(payload, data):
    snapshot = Path('data/kg_sources.json')
    if isinstance(payload, dict) and isinstance(payload.get('data'), dict):
        payload['data'] = data
        out = payload
    else:
        out = data
    snapshot.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding='utf-8')


def _kg_source_match(item, source_id):
    if not isinstance(item, dict):
        return False
    return str(item.get('source_id') or item.get('id') or item.get('name') or '') == str(source_id)


@app.route('/api/kg/sources', methods=['DELETE'])
def kg_sources_delete():
    from datetime import datetime, timezone
    source_id = (request.args.get('source') or '').strip()
    if not source_id:
        return jsonify({"success": False, "error": "Missing source parameter"}), 400
    payload, data = _kg_load_sources_payload()
    if data is None:
        return jsonify({"success": False, "error": "Data source registry missing or invalid"}), 404
    removed_item = None
    for key in ('excel', 'pdf', 'sources'):
        items = data.get(key)
        if not isinstance(items, list):
            continue
        kept = []
        for item in items:
            if _kg_source_match(item, source_id):
                if removed_item is None:
                    removed_item = dict(item)
            else:
                kept.append(item)
        data[key] = kept
    if removed_item is None:
        return jsonify({"success": False, "error": "Data source not found"}), 404
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    trash = [item for item in trash if not _kg_source_match(item, source_id)]
    removed_item['deleted_at'] = datetime.now(timezone.utc).isoformat()
    trash.insert(0, removed_item)
    data['trash'] = trash
    _kg_save_sources_payload(payload, data)
    return jsonify({"success": True, "data": {"deleted": True, "source": source_id, "trash_count": len(trash)}})


@app.route('/api/kg/sources/trash', methods=['GET'])
def kg_sources_trash():
    payload, data = _kg_load_sources_payload()
    if data is None:
        return jsonify({"success": True, "data": {"trash": []}})
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    return jsonify({"success": True, "data": {"trash": trash, "count": len(trash)}})


@app.route('/api/kg/sources/restore', methods=['POST'])
def kg_sources_restore():
    body = request.get_json(silent=True) or {}
    source_id = str(body.get('source') or request.args.get('source') or '').strip()
    if not source_id:
        return jsonify({"success": False, "error": "Missing source parameter"}), 400
    payload, data = _kg_load_sources_payload()
    if data is None:
        return jsonify({"success": False, "error": "Data source registry missing or invalid"}), 404
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    restored = None
    remaining = []
    for item in trash:
        if restored is None and _kg_source_match(item, source_id):
            restored = dict(item)
        else:
            remaining.append(item)
    if restored is None:
        return jsonify({"success": False, "error": "Data source not found in trash"}), 404
    restored.pop('deleted_at', None)
    kind = restored.get('kind') or restored.get('file_type') or 'excel'
    bucket = 'pdf' if str(kind).lower() == 'pdf' else 'excel'
    for key in (bucket, 'sources'):
        items = data.get(key) if isinstance(data.get(key), list) else []
        items = [item for item in items if not _kg_source_match(item, source_id)]
        items.insert(0, restored)
        data[key] = items
    data['trash'] = remaining
    _kg_save_sources_payload(payload, data)
    return jsonify({"success": True, "data": {"restored": True, "source": restored, "trash_count": len(remaining)}})


@app.route('/api/kg/sources/purge', methods=['POST', 'DELETE'])
def kg_sources_purge():
    body = request.get_json(silent=True) or {}
    source_id = str(body.get('source') or request.args.get('source') or '').strip()
    if not source_id:
        return jsonify({"success": False, "error": "Missing source parameter"}), 400
    payload, data = _kg_load_sources_payload()
    if data is None:
        return jsonify({"success": False, "error": "Data source registry missing or invalid"}), 404
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    removed = None
    remaining = []
    for item in trash:
        if removed is None and _kg_source_match(item, source_id):
            removed = dict(item)
        else:
            remaining.append(item)
    if removed is None:
        return jsonify({"success": False, "error": "Data source not found in trash"}), 404
    data['trash'] = remaining
    _kg_save_sources_payload(payload, data)
    return jsonify({"success": True, "data": {"purged": True, "source": source_id, "trash_count": len(remaining)}})


def _kg_skills_path():
    for candidate in (Path('data/kg_skills.json'), Path('7.30_kg/agentic_kg/memory/skill_registry.json')):
        if candidate.exists():
            return candidate
    return Path('data/kg_skills.json')


def _kg_load_skills_payload():
    path = _kg_skills_path()
    if not path.exists():
        return path, {"skills": [], "trash": []}
    try:
        raw = json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return path, {"skills": [], "trash": []}
    if isinstance(raw, list):
        return path, {"skills": raw, "trash": []}
    if isinstance(raw, dict):
        skills = raw.get('skills') if isinstance(raw.get('skills'), list) else []
        trash = raw.get('trash') if isinstance(raw.get('trash'), list) else []
        return path, {"skills": skills, "trash": trash}
    return path, {"skills": [], "trash": []}


def _kg_save_skills_payload(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def _kg_skill_match(item, skill_id):
    if not isinstance(item, dict):
        return False
    sid = str(skill_id or '').strip()
    if not sid:
        return False
    return str(item.get('skill_id') or item.get('id') or '').strip() == sid


def _kg_skills_registry():
    _path, data = _kg_load_skills_payload()
    return data.get('skills') or []


@app.route('/api/kg/skills', methods=['GET'])
def kg_skills():
    _path, data = _kg_load_skills_payload()
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    return jsonify({
        "success": True,
        "data": {
            "skills": data.get('skills') or [],
            "trash_count": len(trash),
        },
    })


@app.route('/api/kg/skills', methods=['DELETE'])
def kg_skills_delete():
    from datetime import datetime, timezone
    skill_id = (request.args.get('skill') or request.args.get('skill_id') or '').strip()
    if not skill_id:
        return jsonify({"success": False, "error": "Missing skill parameter"}), 400
    path, data = _kg_load_skills_payload()
    skills = data.get('skills') if isinstance(data.get('skills'), list) else []
    removed = None
    kept = []
    for item in skills:
        if removed is None and _kg_skill_match(item, skill_id):
            removed = dict(item)
        else:
            kept.append(item)
    if removed is None:
        return jsonify({"success": False, "error": "Skill not found"}), 404
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    trash = [item for item in trash if not _kg_skill_match(item, skill_id)]
    removed['deleted_at'] = datetime.now(timezone.utc).isoformat()
    trash.insert(0, removed)
    data['skills'] = kept
    data['trash'] = trash
    _kg_save_skills_payload(path, data)
    return jsonify({"success": True, "data": {"deleted": True, "skill": skill_id, "trash_count": len(trash)}})


@app.route('/api/kg/skills/status', methods=['POST'])
def kg_skills_status():
    body = request.get_json(silent=True) or {}
    skill_id = str(body.get('skill') or body.get('skill_id') or request.args.get('skill') or '').strip()
    status = str(body.get('status') or '').strip().lower()
    if not skill_id:
        return jsonify({"success": False, "error": "Missing skill parameter"}), 400
    if status not in ('enabled', 'disabled'):
        return jsonify({"success": False, "error": "status must be enabled or disabled"}), 400
    path, data = _kg_load_skills_payload()
    skills = data.get('skills') if isinstance(data.get('skills'), list) else []
    updated = None
    for item in skills:
        if _kg_skill_match(item, skill_id):
            item['status'] = status
            updated = item
            break
    if updated is None:
        return jsonify({"success": False, "error": "Skill not found"}), 404
    data['skills'] = skills
    _kg_save_skills_payload(path, data)
    return jsonify({"success": True, "data": {"skill": updated}})


@app.route('/api/kg/skills/trash', methods=['GET'])
def kg_skills_trash():
    _path, data = _kg_load_skills_payload()
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    return jsonify({"success": True, "data": {"trash": trash, "count": len(trash)}})


@app.route('/api/kg/skills/restore', methods=['POST'])
def kg_skills_restore():
    body = request.get_json(silent=True) or {}
    skill_id = str(body.get('skill') or body.get('skill_id') or request.args.get('skill') or '').strip()
    if not skill_id:
        return jsonify({"success": False, "error": "Missing skill parameter"}), 400
    path, data = _kg_load_skills_payload()
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    restored = None
    remaining = []
    for item in trash:
        if restored is None and _kg_skill_match(item, skill_id):
            restored = dict(item)
        else:
            remaining.append(item)
    if restored is None:
        return jsonify({"success": False, "error": "Skill not found in trash"}), 404
    restored.pop('deleted_at', None)
    skills = data.get('skills') if isinstance(data.get('skills'), list) else []
    skills = [item for item in skills if not _kg_skill_match(item, skill_id)]
    skills.insert(0, restored)
    data['skills'] = skills
    data['trash'] = remaining
    _kg_save_skills_payload(path, data)
    return jsonify({"success": True, "data": {"restored": True, "skill": restored, "trash_count": len(remaining)}})


@app.route('/api/kg/skills/purge', methods=['POST', 'DELETE'])
def kg_skills_purge():
    body = request.get_json(silent=True) or {}
    skill_id = str(body.get('skill') or body.get('skill_id') or request.args.get('skill') or '').strip()
    if not skill_id:
        return jsonify({"success": False, "error": "Missing skill parameter"}), 400
    path, data = _kg_load_skills_payload()
    trash = data.get('trash') if isinstance(data.get('trash'), list) else []
    removed = None
    remaining = []
    for item in trash:
        if removed is None and _kg_skill_match(item, skill_id):
            removed = dict(item)
        else:
            remaining.append(item)
    if removed is None:
        return jsonify({"success": False, "error": "Skill not found in trash"}), 404
    data['trash'] = remaining
    _kg_save_skills_payload(path, data)
    return jsonify({"success": True, "data": {"purged": True, "skill": skill_id, "trash_count": len(remaining)}})


_SKILL_CHAT_SESSIONS = {}


def _skill_chat_slug(text, fallback='custom_skill'):
    raw = re.sub(r'[^a-zA-Z0-9_\u4e00-\u9fff]+', '_', str(text or '').strip()).strip('_').lower()
    if not raw:
        return fallback
    ascii_only = re.sub(r'[^a-z0-9_]+', '', raw)
    return (ascii_only or fallback)[:48]


def _skill_chat_detect(text):
    lower = str(text or '').lower()
    raw = str(text or '')
    file_type = ''
    if re.search(r'\bpdf\b|pdf|扫描件|图纸|规范文档', lower, re.I):
        file_type = 'pdf'
    elif re.search(r'excel|xlsx|xls|csv|表格|sheet|用例表', lower, re.I):
        file_type = 'excel'
    elif re.search(r'增量|diff|变更|版本对比', lower, re.I):
        file_type = 'incremental'
    elif re.search(r'融合|去重|冲突|三元组集合', lower, re.I):
        file_type = 'fusion'

    entities = []
    for label in ['测试类型', '测试用例', '前置条件', '执行步骤', '执行动作', '预期结果', '预期行为', '实体', '关系', '三元组', '条款', '信号', '故障', '灯光', '规则']:
        if label in raw:
            entities.append(label)

    has_extract = bool(re.search(r'抽取|提取|解析|映射|输出|生成|补充|新建', raw))
    has_structure = bool(re.search(r'表头|字段|列|sheet|页|图片|ocr|模板|分组|场景|夜间|转向', lower, re.I))
    skill_kind = ''
    if re.search(r'自动化|自动执行|executable|可执行', raw, re.I):
        skill_kind = 'executable'
    elif re.search(r'人工|指导|instruction|操作说明|手册', raw, re.I):
        skill_kind = 'instruction'

    theme = ''
    theme_match = re.search(
        r'(?:主题[为是：:]\s*|关于\s*|新建\s*skills?[，,：:]*\s*主题[为是：:]\s*)([^\n，。；;]{2,40})',
        raw,
        re.I,
    )
    if theme_match:
        theme = theme_match.group(1).strip()
    elif re.search(r'灯光|抽取规则|补充', raw):
        theme = re.sub(r'^(?:新建\s*skills?[，,\s]*)?', '', raw).strip()[:40]

    name_hint = ''
    name_match = re.search(r'(?:叫作|名为|名称[是为：:]\s*|skill[:：]\s*)([^\n，。；;]{2,40})', raw, re.I)
    if name_match:
        name_hint = name_match.group(1).strip()
    elif theme:
        name_hint = theme[:40]

    return {
        'file_type': file_type,
        'entities': entities,
        'has_extract': has_extract,
        'has_structure': has_structure,
        'skill_kind': skill_kind,
        'theme': theme,
        'name_hint': name_hint,
        'length': len(raw.strip()),
    }


def _skill_chat_merge_profile(session, message):
    detected = _skill_chat_detect(message)
    defaults = {
        'file_type': '',
        'entities': [],
        'has_extract': False,
        'has_structure': False,
        'skill_kind': '',
        'theme': '',
        'name_hint': '',
        'raw_notes': [],
        'clarify_round': 0,
    }
    profile = session.get('profile')
    if not isinstance(profile, dict):
        profile = {}
    for key, value in defaults.items():
        profile.setdefault(key, value if not isinstance(value, list) else list(value))
    session['profile'] = profile
    if detected['file_type']:
        profile['file_type'] = detected['file_type']
    if detected['entities']:
        for item in detected['entities']:
            if item not in profile['entities']:
                profile['entities'].append(item)
    profile['has_extract'] = bool(profile.get('has_extract')) or detected['has_extract']
    profile['has_structure'] = bool(profile.get('has_structure')) or detected['has_structure']
    if detected.get('skill_kind'):
        profile['skill_kind'] = detected['skill_kind']
    if detected.get('theme'):
        profile['theme'] = detected['theme']
    if detected['name_hint']:
        profile['name_hint'] = detected['name_hint']
    note = str(message or '').strip()
    if note and note not in profile['raw_notes']:
        profile['raw_notes'].append(note)
        profile['raw_notes'] = profile['raw_notes'][-8:]
    return profile, detected


def _skill_chat_missing(profile):
    missing = []
    if not profile.get('skill_kind'):
        missing.append('skill_kind')
    if not profile.get('theme') and not profile.get('name_hint') and not profile.get('entities'):
        missing.append('theme')
    if not profile.get('file_type') and not profile.get('theme'):
        missing.append('file_type')
    if not profile.get('entities') and not profile.get('has_extract') and not profile.get('has_structure'):
        missing.append('scope')
    return missing


def _skill_chat_is_lighting(profile):
    notes = ' '.join(profile.get('raw_notes') or [])
    theme = str(profile.get('theme') or profile.get('name_hint') or '')
    blob = theme + ' ' + notes
    return bool(re.search(r'灯光|夜间|转向|抽取规则', blob))


def _skill_chat_build_draft(profile):
    file_type = profile.get('file_type') or 'excel'
    entities = profile.get('entities') or (['灯光', '规则'] if profile.get('theme') else ['测试用例', '关系'])
    name_hint = profile.get('name_hint') or profile.get('theme') or ''
    if not name_hint:
        type_label = {'excel': 'Excel', 'pdf': 'PDF', 'incremental': 'Incremental', 'fusion': 'Fusion'}.get(file_type, file_type.upper())
        name_hint = f'{type_label} {"/".join(entities[:2])} extraction'
    template_map = {
        'excel': 'custom_test_case',
        'pdf': 'pdf_manifest',
        'incremental': 'lock_incremental',
        'fusion': 'triple_set',
    }
    tool_map = {
        'excel': 'excel.extract_custom',
        'pdf': 'pdf.run_online_pipeline',
        'incremental': 'incremental.case_scoped_diff',
        'fusion': 'pdf_kg_pipeline.step5_fuse',
    }
    template_kind = template_map.get(file_type, 'custom')
    skill_id = f"{file_type}.{_skill_chat_slug(template_kind)}.{_skill_chat_slug(name_hint, 'v1')}"
    if not skill_id.endswith('.v1'):
        skill_id = skill_id.rstrip('.') + '.v1'
    kind = profile.get('skill_kind') or 'executable'

    if _skill_chat_is_lighting(profile):
        return {
            'name': 'Lighting-scene test case extraction rule supplement',
            'function_desc': 'Extend automated extraction for night and turn-signal lighting test cases',
            'audience': 'Grouped Excel test case files with night/turn-signal lighting scenarios',
            'input_req': (
                'Excel input with at least one sheet for grouping (ID, name, title); '
                'test data on the second sheet with a complex Chinese header row; must include night or turn-signal lighting items'
            ),
            'output_result': (
                'Structured night/turn-signal lighting test case entries; '
                'each entry has Preconditions, Actions, Expectations'
            ),
            'scenario': 'Automated extraction for new lighting test scenarios in automotive manufacturing',
            'steps': (
                '1.Use the second Excel sheet as the data source\n'
                '2.Parse the complex Chinese header; group by ID, name, title\n'
                '3.Filter groups containing night or turn-signal lighting keywords\n'
                '4.Extract Preconditions, Actions, Expectations per group\n'
            ),
            'limits': (
                'Only supports Excel files with explicit ID-based grouping\n'
                'Non-standard lighting terms require a unified naming convention first'
            ),
            'example': (
                'For a group with auto low-beam at night and left turn-signal blink, '
                'the system extracts preconditions, steps, and expected results as two structured cases'
            ),
            'skill_id': skill_id[:80],
            'file_type': file_type,
            'template_kind': template_kind,
            'tool': tool_map.get(file_type, 'custom.extract'),
            'skill_kind': kind,
            'description': 'Automated extraction for night and turn-signal lighting test cases',
        }

    description = '; '.join(profile.get('raw_notes') or [])[:220] or f'Extract {"、".join(entities)} from {file_type}'
    return {
        'name': name_hint[:40],
        'function_desc': description,
        'audience': f'For the specified {file_type} data source',
        'input_req': f'Input type {file_type} with recognizable headers or page structure',
        'output_result': f'Structured {"、".join(entities)} entries',
        'scenario': 'Extract per confirmed scope; output candidates for human review',
        'steps': '1.Detect input structure\n2.Map fields to entities/relations\n3.Output structured rows and flag review items',
        'limits': 'Only for declared source shapes\nMissing fields flagged for review without stopping the batch',
        'example': 'Extraction output follows user-provided topic and scope',
        'skill_id': skill_id[:80],
        'file_type': file_type,
        'template_kind': template_kind,
        'tool': tool_map.get(file_type, 'custom.extract'),
        'skill_kind': kind,
        'description': description,
    }


def _skill_chat_clarify_prompt(profile, missing):
    questions = []
    notes = profile.get('raw_notes') or []
    lighting_theme = bool(profile.get('theme')) or any(
        ('灯光' in n) or ('抽取规则' in n) or ('补充' in n) for n in notes
    )
    if 'skill_kind' in missing or not profile.get('skill_kind'):
        questions.append('1.Is this Skill for human guidance or automation (instruction vs executable)?')
    if lighting_theme or 'theme' in missing or 'scope' in missing or 'file_type' in missing or not profile.get('has_structure'):
        idx = len(questions) + 1
        if lighting_theme:
            questions.append(
                f'{idx}.What lighting extraction rules should be added—fix, extend, or new scenes (e.g. night, turn signal)?'
            )
        else:
            questions.append(
                f'{idx}.Specify file type (Excel/PDF), entities/relations to extract, or scope of the rule change.'
            )
    if not questions:
        questions.append('Add file type, extraction targets, or scenario so a draft can be generated.')
    return 'Please provide:\n' + '\n'.join(questions)


def _skill_chat_reply(session, message):
    profile, detected = _skill_chat_merge_profile(session, message)
    text = str(message or '').strip()
    missing = _skill_chat_missing(profile)
    profile['clarify_round'] = int(profile.get('clarify_round') or 0) + 1

    if not text or detected['length'] < 2:
        return {
            'reply': 'What file type should this Skill handle, and which entities/relations to extract?',
            'draft': None,
            'stage': 'clarify',
        }

    if re.search(r'^(你好|您好|在吗|hello|hi)[\s!！。.?？]*$', text, re.I):
        return {
            'reply': 'Hello. Describe the Skill to create or change: topic, file type, extraction goals.',
            'draft': None,
            'stage': 'clarify',
        }

    ready_for_draft = bool(profile.get('skill_kind')) and bool(
        profile.get('theme')
        or profile.get('name_hint')
        or profile.get('entities')
        or profile.get('has_structure')
        or (profile.get('has_extract') and profile.get('file_type'))
    )

    if missing and not ready_for_draft:
        return {
            'reply': _skill_chat_clarify_prompt(profile, missing),
            'draft': None,
            'stage': 'clarify',
        }

    if not profile.get('file_type'):
        profile['file_type'] = 'excel'
    draft = _skill_chat_build_draft(profile)
    session['draft'] = draft
    return {
        'reply': 'Draft ready for confirmation',
        'draft': draft,
        'stage': 'draft',
    }


@app.route('/api/kg/skills/pending', methods=['GET'])
def kg_skills_pending():
    pending_dir = Path('data/kg_skill_drafts')
    drafts = []
    if pending_dir.exists():
        for path in sorted(pending_dir.glob('*.json'), key=lambda p: p.stat().st_mtime, reverse=True):
            try:
                payload = json.loads(path.read_text(encoding='utf-8'))
            except Exception:
                payload = {'name': path.stem}
            drafts.append({
                'filename': path.name,
                'name': payload.get('name') or path.stem,
                'skill_id': payload.get('skill_id') or '',
                'file_type': payload.get('file_type') or '',
                'template_kind': payload.get('template_kind') or '',
                'description': payload.get('function_desc') or payload.get('description') or '',
                'function_desc': payload.get('function_desc') or '',
                'status': payload.get('status') or 'pending_review',
                'created_at': payload.get('created_at') or '',
            })
    return jsonify({"success": True, "data": {"drafts": drafts, "pending": drafts}})


@app.route('/api/kg/skills/chat', methods=['POST'])
def kg_skills_chat():
    payload = request.get_json(silent=True) or {}
    conversation_id = str(payload.get('conversation_id') or '').strip() or f"skill-{uuid.uuid4().hex[:12]}"
    message = ''
    history = []
    if isinstance(payload.get('message'), str):
        message = payload['message'].strip()
    if isinstance(payload.get('messages'), list):
        history = [item for item in payload['messages'] if isinstance(item, dict)]
        if not message and history:
            last = history[-1] or {}
            message = str(last.get('content') or '').strip()

    session = _SKILL_CHAT_SESSIONS.setdefault(conversation_id, {'profile': {}, 'messages': []})
    if history and not session.get('messages'):
        for item in history[:-1]:
            content = str(item.get('content') or '').strip()
            if content and item.get('role') == 'user':
                _skill_chat_merge_profile(session, content)
    result = _skill_chat_reply(session, message)
    session['messages'] = (session.get('messages') or []) + [
        {'role': 'user', 'content': message},
        {'role': 'assistant', 'content': result['reply']},
    ]
    return jsonify({
        "success": True,
        "data": {
            "conversation_id": conversation_id,
            "reply": result['reply'],
            "draft": result.get('draft'),
            "stage": result.get('stage') or 'clarify',
        },
    })


@app.route('/api/kg/skills/draft', methods=['POST', 'DELETE'])
def kg_skills_draft():
    pending_dir = Path('data/kg_skill_drafts')
    pending_dir.mkdir(parents=True, exist_ok=True)
    if request.method == 'DELETE':
        filename = (request.args.get('filename') or '').strip()
        if not filename:
            return jsonify({"success": False, "error": "Missing filename"}), 400
        target = pending_dir / Path(filename).name
        if target.exists():
            target.unlink()
        return jsonify({"success": True, "data": {"deleted": True, "filename": filename}})

    payload = request.get_json(silent=True) or {}
    if not payload.get('confirmed'):
        return jsonify({"success": False, "error": "Confirm the draft before creating"}), 400
    name = str(payload.get('name') or 'untitled_skill').strip() or 'untitled_skill'
    skill_id = str(payload.get('skill_id') or _skill_chat_slug(name)).strip()
    filename = f"{_skill_chat_slug(skill_id or name)}.draft.json"
    record = dict(payload)
    record.update({
        'name': name,
        'skill_id': skill_id,
        'status': 'pending_review',
        'created_at': datetime.now(timezone.utc).isoformat(),
    })
    (pending_dir / filename).write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    return jsonify({"success": True, "data": {"filename": filename, "name": name, "status": "pending_review", "skill_id": skill_id}})


@app.route('/api/kg/incr/state', methods=['GET'])
def kg_incr_state():
    cases = []
    try:
        cases = _load_search_cases()
    except RuntimeError:
        cases = []
    return jsonify({
        "success": True,
        "data": {
            "has_baseline": bool(cases),
            "domain": request.args.get('domain') or 'Test cases',
            "current": {"cases": len(cases), "triples": len(cases) * 3},
        },
    })


def _kg_find_source(source_id):
    if not source_id:
        return None
    _, data = _kg_load_sources_payload()
    if not data:
        return None
    for key in ('sources', 'excel', 'pdf'):
        items = data.get(key)
        if not isinstance(items, list):
            continue
        for item in items:
            if _kg_source_match(item, source_id):
                return dict(item)
    return None


def _kg_preview_excel_cases(limit=40):
    cases = []
    try:
        loaded = _load_search_cases()
    except RuntimeError:
        loaded = []
    for index, item in enumerate(loaded[:limit]):
        cases.append({
            "row": index,
            "case_id": item.get('case_id'),
            "case_name": item.get('case_name'),
            "case_type": item.get('test_type'),
            "preconditions": item.get('preconditions') or [],
            "step": '；'.join(item.get('actions') or []),
            "expectation": '；'.join(item.get('expected_behaviors') or []),
        })
    return cases


def _kg_preview_pdf_triples(source, limit=40):
    name = (source or {}).get('name') or (source or {}).get('source_id') or 'PDF'
    total = int((source or {}).get('triple_count') or (source or {}).get('triples') or 0)
    count = min(limit, total if total > 0 else 24)
    triples = []
    for index in range(count):
        triples.append({
            "head": f"{name} · segment {index + 1}",
            "relation": "描述",
            "tail": f"Preview relation {index + 1}",
            "keep": True,
        })
    return triples


@app.route('/api/kg/build', methods=['POST'])
def kg_build():
    from datetime import datetime, timezone
    import uuid

    upload = request.files.get('file')
    if upload is None or not (upload.filename or '').strip():
        return jsonify({"success": False, "error": "Please select an Excel, CSV, or PDF file"}), 400

    filename = (upload.filename or '').strip()
    domain = (request.form.get('domain') or '').strip()
    mode = (request.form.get('mode') or 'supplement').strip() or 'supplement'
    lower = filename.lower()
    if lower.endswith('.pdf'):
        kind = 'pdf'
    elif lower.endswith('.csv'):
        kind = 'csv'
    else:
        kind = 'excel'

    try:
        upload.stream.seek(0, 2)
        size = int(upload.stream.tell() or 0)
        upload.stream.seek(0)
    except Exception:
        size = 0

    source_id = 'build-' + uuid.uuid4().hex[:16]
    now = datetime.now(timezone.utc).isoformat()
    cases = []
    triples = []
    if kind == 'pdf':
        triples = _kg_preview_pdf_triples({
            "name": filename,
            "source_id": source_id,
            "triple_count": 24,
        })
        triples_total = len(triples)
        case_count = 0
    else:
        cases = _kg_preview_excel_cases()
        case_count = len(cases)
        triples_total = case_count * 3

    pending = {
        "add": {"count": case_count or triples_total, "records": []},
        "remove": {"count": 0, "records": []},
        "changed": {"count": 0, "records": []},
    }

    return jsonify({
        "success": True,
        "data": {
            "source_id": source_id,
            "job_id": 'job-' + source_id,
            "kind": kind,
            "file_name": filename,
            "name": filename,
            "domain": domain,
            "mode": mode,
            "status": "ready",
            "version": 1,
            "bytes": size,
            "imported": False,
            "duplicate": False,
            "triples_total": int(triples_total or 0),
            "case_count": int(case_count or 0),
            "cases": cases,
            "triples": triples,
            "pending": pending,
            "updated_at": now,
            "message": "Parsing complete; confirm to submit for review.",
        },
    })


@app.route('/api/kg/preview', methods=['GET'])
def kg_preview():
    source_id = (request.args.get('source') or request.args.get('source_id') or '').strip()
    mode = (request.args.get('mode') or 'supplement').strip() or 'supplement'
    source = _kg_find_source(source_id)
    if not source and not source_id:
        return jsonify({"success": False, "error": "Missing source parameter"}), 400
    if not source:
        source = {
            "source_id": source_id,
            "id": source_id,
            "name": source_id,
            "kind": "excel",
            "status": "ready",
            "triple_count": 0,
            "case_count": 0,
        }

    kind = str(source.get('kind') or source.get('file_type') or 'excel').lower()
    if kind not in ('excel', 'pdf', 'csv'):
        kind = 'excel'
    status_name = str(source.get('status') or 'ready')
    triples_total = source.get('triple_count')
    if triples_total is None:
        triples_total = source.get('triples') or 0
    case_count = source.get('case_count') or 0
    imported = status_name in ('imported', 'published', 'available', 'ready', 'reviewed', 'completed')

    cases = []
    triples = []
    if kind == 'pdf':
        triples = _kg_preview_pdf_triples(source)
    else:
        cases = _kg_preview_excel_cases()
        if not case_count:
            case_count = len(cases)

    pending = {
        "add": {"count": 0, "records": []},
        "remove": {"count": 0, "records": []},
        "changed": {"count": 0, "records": []},
    }

    return jsonify({
        "success": True,
        "data": {
            "source_id": source.get('source_id') or source.get('id') or source_id,
            "kind": kind,
            "file_name": source.get('name') or source.get('file_name') or source_id,
            "name": source.get('name') or source.get('file_name') or source_id,
            "domain": source.get('domain') or '',
            "mode": mode,
            "status": status_name if status_name not in ('available',) else 'ready',
            "version": source.get('version'),
            "imported": False,
            "duplicate": imported,
            "triples_total": int(triples_total or len(triples) or 0),
            "case_count": int(case_count or len(cases) or 0),
            "cases": cases,
            "triples": triples,
            "pending": pending,
            "summary": {
                "cases": int(case_count or len(cases) or 0),
                "triples": int(triples_total or len(triples) or 0),
            },
            "preview_tables": [],
            "message": "This file already exists. Parse results loaded." if imported else "Parsing complete; confirm to import.",
        },
    })


@app.route('/api/kg/import', methods=['POST'])
def kg_import():
    from datetime import datetime, timezone
    import uuid

    body = request.get_json(silent=True) or {}
    source_id = str(body.get('source_id') or body.get('source') or '').strip()
    if not source_id:
        source_id = 'build-' + uuid.uuid4().hex[:16]

    filename = str(body.get('file_name') or body.get('name') or source_id).strip() or source_id
    domain = str(body.get('domain') or '').strip()
    mode = str(body.get('mode') or 'supplement').strip() or 'supplement'
    kind = str(body.get('kind') or 'excel').lower()
    if kind not in ('excel', 'pdf', 'csv'):
        kind = 'pdf' if filename.lower().endswith('.pdf') else 'excel'

    cases = body.get('cases') if isinstance(body.get('cases'), list) else []
    triples = body.get('triples') if isinstance(body.get('triples'), list) else []
    case_count = body.get('case_count')
    if case_count is None:
        case_count = len(cases)
    triples_total = body.get('triples_total')
    if triples_total is None:
        triples_total = body.get('triple_count')
    if triples_total is None:
        triples_total = len(triples) if triples else (int(case_count or 0) * 3)

    existing = _kg_find_source(source_id)
    existing_status = str((existing or {}).get('status') or '')
    # When submitting from the library for review, keep the original entry and create a pending copy
    preserve_library = bool(existing and existing_status and existing_status != 'pending_review')
    pending_id = ('build-' + uuid.uuid4().hex[:16]) if preserve_library else source_id
    base_source = existing or {}
    if not domain:
        domain = str(base_source.get('domain') or '').strip()
    if not filename or filename == source_id:
        filename = str(base_source.get('name') or base_source.get('file_name') or filename).strip() or pending_id
    if kind == 'excel' and str(base_source.get('kind') or '').lower() in ('pdf', 'csv', 'excel'):
        kind = str(base_source.get('kind') or kind).lower()
    if case_count is None or int(case_count or 0) == 0:
        case_count = base_source.get('case_count') or case_count or 0
    if triples_total is None or int(triples_total or 0) == 0:
        triples_total = base_source.get('triple_count') or base_source.get('triples') or triples_total or 0

    now = datetime.now(timezone.utc).isoformat()
    entry = {
        "id": pending_id,
        "source_id": pending_id,
        "kind": kind,
        "name": filename,
        "file_name": filename,
        "domain": domain,
        "mode": mode,
        "status": "pending_review",
        "case_count": int(case_count or 0),
        "triple_count": int(triples_total or 0),
        "triples": int(triples_total or 0),
        "version": int(body.get('version') or base_source.get('version') or 1),
        "bytes": int(body.get('bytes') or base_source.get('bytes') or 0),
        "updated_at": now,
        "job_id": body.get('job_id') or ('job-' + pending_id),
    }
    if preserve_library:
        entry['from_source_id'] = source_id

    payload, data = _kg_load_sources_payload()
    if data is None:
        data = {"excel": [], "pdf": [], "sources": [], "trash": [], "graph": {}}
        payload = {"data": data}

    remove_id = None if preserve_library else source_id
    if remove_id:
        for key in ('excel', 'pdf', 'sources'):
            items = data.get(key)
            if not isinstance(items, list):
                data[key] = []
                continue
            data[key] = [item for item in items if not _kg_source_match(item, remove_id)]

    bucket = 'pdf' if kind == 'pdf' else 'excel'
    data.setdefault(bucket, []).insert(0, entry)
    if not isinstance(data.get('sources'), list):
        data['sources'] = []
    if remove_id:
        data['sources'] = [item for item in data['sources'] if not _kg_source_match(item, remove_id)]
    data['sources'].insert(0, dict(entry))

    _kg_save_sources_payload(payload, data)
    return jsonify({
        "success": True,
        "data": {
            "source_id": pending_id,
            "status": "pending_review",
            "imported": False,
            "pending_review": True,
            "total": int(triples_total or 0),
            "triples_total": int(triples_total or 0),
            "triple_count": int(triples_total or 0),
            "case_count": int(case_count or 0),
            "source": entry,
            "message": "Submitted for review",
        },
    })


@app.route('/api/kg/jobs/<job_id>', methods=['GET'])
@app.route('/api/kg/jobs/<job_id>/retry', methods=['POST'])
def kg_jobs(job_id):
    return jsonify({"success": True, "data": {"job_id": job_id, "status": "failed", "error": "No local build job", "result": {}}})


_EXCEL_DOC_CACHE = {"key": None, "mtime": None, "doc": None}


def _kg_split_cell_items(value):
    text = '' if value is None else str(value).strip()
    if not text or text.lower() == 'nan':
        return []
    parts = re.split(r'[\n；;]+', text)
    items = []
    for part in parts:
        cleaned = re.sub(r'^\s*\d+[\.、\)]\s*', '', part).strip()
        if cleaned:
            items.append(cleaned)
    return items or ([text] if text else [])


def _kg_excel_path_for_source(source):
    name = str((source or {}).get('name') or (source or {}).get('file_name') or '').strip()
    data_dir = Path('data')
    if name:
        direct = data_dir / name
        if direct.exists():
            return direct
    for path in sorted(data_dir.glob('*.xlsx')):
        if name and (path.name == name or name in path.name or path.stem in name):
            return path
    return None


def _kg_build_excel_case_triples(case_name, case_type, preconditions, steps, expectations):
    title = case_name or 'Unnamed case'
    triples = []
    if case_type:
        triples.append({"head": title, "relation": "属于类型", "tail": case_type, "keep": True})
    for item in preconditions:
        triples.append({"head": title, "relation": "前提条件", "tail": item, "keep": True})
    for item in steps:
        triples.append({"head": title, "relation": "操作动作", "tail": item, "keep": True})
    for item in expectations:
        triples.append({"head": title, "relation": "预期结果", "tail": item, "keep": True})
    if not triples:
        triples.append({"head": title, "relation": "待补充", "tail": "No extractable content", "keep": True})
    return triples


def _kg_parse_excel_workbook(path):
    import pandas as pd

    frames = []
    xl = pd.ExcelFile(path)
    for sheet in xl.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet, header=None)
        if df is None or df.empty:
            continue
        frames.append(df)
        # Prefer the first non-empty sheet that looks like the case table.
        header = [str(cell).strip() if cell is not None and str(cell) != 'nan' else '' for cell in df.iloc[0].tolist()]
        if any(name in header for name in ('用例名称', '测试步骤', '期望结果', '前置条件', '测试功能')):
            break

    if not frames:
        return []

    df = frames[-1]
    header = [str(cell).strip() if cell is not None and str(cell) != 'nan' else '' for cell in df.iloc[0].tolist()]

    def col(*names):
        for name in names:
            if name in header:
                return header.index(name)
        return None

    idx_type = col('测试功能', '测试用例类型', '用例类型')
    idx_name = col('用例名称', '测试用例名称', '标题')
    idx_pre = col('前置条件', '前提条件')
    idx_step = col('测试步骤', '执行动作', '步骤')
    idx_exp = col('期望结果', '预期行为', '预期结果')
    idx_id = col('编号', '用例编号', 'ID')

    cases = []
    for offset, (_, series) in enumerate(df.iloc[1:].iterrows(), start=1):
        values = series.tolist()

        def cell(index):
            if index is None or index >= len(values):
                return ''
            value = values[index]
            if value is None:
                return ''
            text = str(value).strip()
            return '' if text.lower() == 'nan' else text

        case_name = cell(idx_name)
        case_type = cell(idx_type)
        if not case_name and not case_type:
            continue
        preconditions = _kg_split_cell_items(cell(idx_pre))
        steps = _kg_split_cell_items(cell(idx_step))
        expectations = _kg_split_cell_items(cell(idx_exp))
        case_id = cell(idx_id) or f'ROW_{offset}'
        display_name = case_name or case_id
        triples = _kg_build_excel_case_triples(display_name, case_type, preconditions, steps, expectations)
        cases.append({
            "row": offset,
            "case_id": case_id,
            "case_name": display_name,
            "case_type": case_type,
            "preconditions": preconditions,
            "step": '；'.join(steps),
            "expectation": '；'.join(expectations),
            "triples": triples,
            "summary": "",
            "status": "pending_review",
            "reviewed": False,
            "review_version": 0,
            "history": [],
            "_dirty": False,
        })
    return cases


def _kg_excel_cases_from_testcase_v2(source):
    name = str((source or {}).get('name') or '')
    try:
        loaded = _load_search_cases()
    except RuntimeError:
        loaded = []
    filtered = loaded
    if '外灯' in name:
        filtered = [
            item for item in loaded
            if '外灯' in (item.get('test_type') or '') or '外灯' in (item.get('case_name') or '')
        ]
    cases = []
    for offset, item in enumerate(filtered, start=1):
        preconditions = item.get('preconditions') or []
        steps = item.get('actions') or []
        expectations = item.get('expected_behaviors') or []
        case_name = item.get('case_name') or item.get('case_id')
        case_type = item.get('test_type') or ''
        cases.append({
            "row": offset,
            "case_id": item.get('case_id'),
            "case_name": case_name,
            "case_type": case_type,
            "preconditions": preconditions,
            "step": '；'.join(steps),
            "expectation": '；'.join(expectations),
            "triples": _kg_build_excel_case_triples(case_name, case_type, preconditions, steps, expectations),
            "summary": "",
            "status": "pending_review",
            "reviewed": False,
            "review_version": 0,
            "history": [],
            "_dirty": False,
        })
    return cases


def _kg_load_excel_doc(source_id):
    source = _kg_find_source(source_id) or {
        "source_id": source_id,
        "id": source_id,
        "name": source_id or 'Excel',
        "kind": "excel",
        "version": 1,
        "domain": "",
    }
    path = _kg_excel_path_for_source(source)
    cache_key = str(path.resolve()) if path else f'v2::{source_id}'
    mtime = path.stat().st_mtime if path and path.exists() else Path('data/TestCase_v2.json').stat().st_mtime if Path('data/TestCase_v2.json').exists() else 0
    cached = _EXCEL_DOC_CACHE
    if cached["key"] == cache_key and cached["mtime"] == mtime and cached["doc"]:
        return cached["doc"]

    if path and path.exists():
        cases = _kg_parse_excel_workbook(path)
        file_name = path.name
    else:
        cases = _kg_excel_cases_from_testcase_v2(source)
        file_name = source.get('name') or source_id or 'Excel'

    domain = source.get('domain') or Path(file_name).stem or 'Test cases'
    doc = {
        "source": {
            "id": source.get('source_id') or source.get('id') or source_id,
            "source_id": source.get('source_id') or source.get('id') or source_id,
            "name": file_name,
            "kind": "excel",
            "version": source.get('version') or 1,
            "domain": domain,
        },
        "domain": domain,
        "version": source.get('version') or 1,
        "file_name": file_name,
        "matched_skill": "excel.grouped_test_case.v1",
        "cases": cases,
    }
    _EXCEL_DOC_CACHE["key"] = cache_key
    _EXCEL_DOC_CACHE["mtime"] = mtime
    _EXCEL_DOC_CACHE["doc"] = doc
    return doc


@app.route('/api/kg/excel/rows', methods=['GET'])
def kg_excel_rows():
    source = (request.args.get('source') or '').strip()
    doc = _kg_load_excel_doc(source)
    return jsonify({"success": True, "data": doc})


@app.route('/api/kg/excel/row/save', methods=['POST'])
def kg_excel_row_save():
    return jsonify({"success": True, "data": {"saved": True, "review_version": 1}})


@app.route('/api/kg/excel/row/reextract', methods=['POST'])
def kg_excel_row_reextract():
    payload = request.get_json(silent=True) or {}
    existing = payload.get('existing_triples') or []
    triples = [{"head": row[0], "relation": row[1], "tail": row[2], "keep": True} for row in existing if isinstance(row, (list, tuple)) and len(row) >= 3]
    return jsonify({"success": True, "data": {"triples": triples, "summary": "Local env returned existing triples (re-extract not connected)"}})


@app.route('/api/kg/excel/graph', methods=['GET'])
def kg_excel_graph():
    source = (request.args.get('source') or '').strip()
    doc = _kg_load_excel_doc(source)
    edges = []
    nodes = set()
    for case in doc.get('cases') or []:
        for triple in case.get('triples') or []:
            if triple.get('keep') is False:
                continue
            head = triple.get('head') or ''
            tail = triple.get('tail') or ''
            relation = triple.get('relation') or ''
            if not head or not tail:
                continue
            edges.append({"head": head, "relation": relation, "tail": tail, "keep": True})
            nodes.add(head)
            nodes.add(tail)
            if len(edges) >= 1200:
                break
        if len(edges) >= 1200:
            break
    return jsonify({
        "success": True,
        "data": {
            "edges": edges,
            "nodes": [{"id": name, "label": name} for name in sorted(nodes)],
        },
    })


@app.route('/api/kg/excel/export', methods=['POST'])
def kg_excel_export():
    return jsonify({"success": False, "error": "Export not available in local preview"}), 501


@app.route('/api/kg/pdf/images', methods=['GET'])
def kg_pdf_images():
    return jsonify({"success": True, "data": {"images": [], "doc": request.args.get('doc'), "source": {"id": request.args.get('doc'), "kind": "pdf", "version": 1}}})


@app.route('/api/kg/pdf/image', methods=['GET'])
def kg_pdf_image():
    return jsonify({"success": True, "data": {
        "doc": request.args.get('doc'),
        "img": request.args.get('img'),
        "triples": [],
        "ocr_text": "",
        "review_version": 0,
        "history": [],
        "source": {"id": request.args.get('doc'), "kind": "pdf", "version": 1},
    }})


@app.route('/api/kg/pdf/eval', methods=['POST'])
def kg_pdf_eval():
    return jsonify({"success": False, "error": "PDF evaluation not available in local preview"}), 501


@app.route('/api/kg/pdf/apply', methods=['POST'])
def kg_pdf_apply():
    return jsonify({"success": False, "error": "PDF apply not available in local preview"}), 501


@app.route('/api/kg/pdf/restore', methods=['POST'])
def kg_pdf_restore():
    return jsonify({"success": False, "error": "PDF restore not available in local preview"}), 501


@app.route('/api/kg/pdf/test-cases', methods=['GET'])
@app.route('/api/kg/pdf/test-cases/preview', methods=['POST'])
def kg_pdf_test_cases():
    return jsonify({"success": True, "data": {"cases": [], "scenarios": []}})


def _qa_question_tokens(question):
    text = re.sub(r'[^\w\u4e00-\u9fff]+', ' ', (question or '').lower(), flags=re.UNICODE).strip()
    parts = [part for part in text.split() if len(part) >= 2]
    tokens = []
    for part in parts or ([text.replace(' ', '')] if text else []):
        if re.search(r'[\u4e00-\u9fff]', part):
            for size in (6, 5, 4, 3, 2):
                if len(part) < size:
                    continue
                for index in range(0, len(part) - size + 1):
                    tokens.append(part[index:index + size])
        else:
            tokens.append(part)
    tokens.sort(key=len, reverse=True)
    unique = []
    seen = set()
    for token in tokens:
        if token in seen:
            continue
        seen.add(token)
        unique.append(token)
    return unique[:40]


def _qa_retrieve_cases(question, limit=8):
    cases = _load_search_cases()
    tokens = _qa_question_tokens(question)
    ask_types = any(key in (question or '') for key in ('测试类型', '有哪些类型', '类型有哪些'))

    if ask_types:
        by_type = {}
        for item in cases:
            test_type = item.get('test_type') or ''
            if not test_type or test_type in by_type:
                continue
            row = {key: value for key, value in item.items() if key != '_blob'}
            row['score'] = 0.9
            row['case_type'] = test_type
            row['source'] = row.get('source') or 'TestCase_v2.json'
            by_type[test_type] = row
            if len(by_type) >= limit:
                break
        return list(by_type.values())

    scored = []
    for item in cases:
        blob = item.get('_blob') or ''
        if not tokens:
            continue
        hits = 0
        weight = 0.0
        for token in tokens:
            if token not in blob:
                continue
            hits += 1
            weight += max(len(token) / 8.0, 0.15)
        if not hits:
            continue
        score = min(weight, 1.2)
        name = (item.get('case_name') or '').lower()
        test_type = (item.get('test_type') or '').lower()
        if any(token in name for token in tokens[:12]):
            score += 0.35
        if any(token in test_type for token in tokens[:12]):
            score += 0.2
        row = {key: value for key, value in item.items() if key != '_blob'}
        row['score'] = round(min(score, 0.999), 4)
        row['case_type'] = row.get('test_type') or ''
        row['source'] = row.get('source') or 'TestCase_v2.json'
        scored.append(row)

    scored.sort(key=lambda row: (-row.get('score', 0), row.get('case_name') or ''))
    return scored[:limit]


def _qa_build_answer(question, cases):
    if not cases:
        return (
            "Not enough relevant test case evidence was retrieved from the knowledge graph.\n\n"
            "Try more specific test types, case names, or feature keywords."
        )

    type_names = []
    seen_types = set()
    for item in cases:
        test_type = item.get('test_type') or item.get('case_type') or ''
        if test_type and test_type not in seen_types:
            seen_types.add(test_type)
            type_names.append(test_type)

    lines = ["Reference answer below.", "", "**Supported by evidence**:"]
    if '测试类型' in (question or '') or '有哪些' in (question or ''):
        lines.append("From retrieved test case evidence, the knowledge graph includes these test types:")
        for index, name in enumerate(type_names[:12], start=1):
            related = next((item for item in cases if (item.get('test_type') or item.get('case_type')) == name), None)
            citation = f"[{cases.index(related) + 1}]" if related in cases else f"[{index}]"
            lines.append(f"- **{name}**{citation}")
    else:
        lines.append("From retrieved evidence, these cases are relevant:")
        for index, item in enumerate(cases[:6], start=1):
            name = item.get('case_name') or item.get('case_id') or 'Unnamed case'
            test_type = item.get('test_type') or item.get('case_type') or '未分类'
            lines.append(f"- **{name}** ({test_type}) [{index}]")

    sample = cases[0]
    preconditions = sample.get('preconditions') or []
    actions = sample.get('actions') or []
    expected = sample.get('expected_behaviors') or []
    lines.extend([
        "",
        "**Sample case structure**:",
        f"- **Case**: {sample.get('case_name') or sample.get('case_id')}[1]",
        f"- **Type**: {sample.get('test_type') or sample.get('case_type') or '未分类'}",
        f"- **Preconditions**: {'; '.join(preconditions) if preconditions else 'Not provided'}",
        f"- **Actions**: {'; '.join(actions) if actions else 'Not provided'}",
        f"- **Expected behavior**: {'; '.join(expected) if expected else 'Not provided'}",
        "",
        "**Note**: Answer from local case retrieval; open evidence on the right for details.",
    ])
    return '\n'.join(lines)


@app.route('/api/qa/chat', methods=['POST'])
def qa_chat():
    """Answer a knowledge question with retrieved case evidence."""
    started = time.time()
    payload = request.get_json(silent=True) or {}
    messages = payload.get('messages') or []
    question = ''
    for message in reversed(messages):
        if isinstance(message, dict) and message.get('role') == 'user' and message.get('content'):
            question = str(message.get('content')).strip()
            break
    if not question:
        question = str(payload.get('question') or payload.get('q') or '').strip()
    if not question:
        return jsonify({"success": False, "error": "Please enter a question"}), 400

    try:
        cases = _qa_retrieve_cases(question, limit=8)
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 500

    answer = _qa_build_answer(question, cases)
    elapsed = time.time() - started
    return jsonify({
        "success": True,
        "data": {
            "answer": answer,
            "elapsed_time": round(elapsed, 3),
            "generation_degraded": True,
            "generation_error_kind": "local-retrieval",
            "evidence": {
                "cases": cases,
                "triples": [],
                "pdf": [],
            },
        },
        "source": "testcase-dataset",
    })


@app.route('/api/pipeline/config', methods=['GET'])
def pipeline_config():
    """Frontend orchestration config for independently deployed pipeline stages."""
    return jsonify({
        "success": True,
        "data": {
            "kg_workbench_url": get_kg_workbench_url(),
        }
    })

@app.route('/sts-generation')
def sts_generation():
    """STS PDF driven test-case generation page."""
    return render_template('sts_generation.html')

@app.route('/skills-platform')
def skills_platform():
    """Skill authoring and marketplace page."""
    return render_template('skills_platform.html')


@app.route('/api/skills', methods=['GET'])
def list_skills():
    return jsonify({"success": True, "data": skill_registry.list()})


@app.route('/api/skills/reload', methods=['POST'])
def reload_skills():
    skill_registry.reload()
    return jsonify({"success": True, "data": skill_registry.list()})


@app.route('/api/skills', methods=['POST'])
def create_skill():
    data = request.get_json(silent=True) or {}
    try:
        skill = skill_registry.create_prompt_skill(
            name=data.get("name", ""),
            description=data.get("description", ""),
            prompt=data.get("prompt", ""),
            model=data.get("model", "Qwen3-4B"),
        )
        return jsonify({"success": True, "data": skill})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400


@app.route('/api/skills/<path:skill_name>/run', methods=['POST'])
def run_skill(skill_name):
    data = request.get_json(silent=True) or {}
    try:
        result = skill_registry.run(skill_name, data)
        return jsonify({"success": True, "data": result})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400


@app.route('/api/skills/search-online', methods=['POST'])
def search_online_skills():
    data = request.get_json(silent=True) or {}
    try:
        result = skill_registry.search_online(data.get("query", ""))
        return jsonify({"success": True, "data": result})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/skills/import-online', methods=['POST'])
def import_online_skill():
    data = request.get_json(silent=True) or {}
    try:
        skill = skill_registry.import_online_skill(data)
        return jsonify({"success": True, "data": skill})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400


@app.route('/api/sts/upload', methods=['POST'])
def upload_sts_pdf():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"success": False, "error": "Missing PDF file"}), 400
    if not file.filename.lower().endswith(".pdf"):
        return jsonify({"success": False, "error": "Only PDF files are supported"}), 400

    document_id = f"sts_{uuid.uuid4().hex[:12]}"
    filename = secure_filename(file.filename) or f"{document_id}.pdf"
    pdf_path = UPLOAD_DIR / f"{document_id}_{filename}"
    file.save(str(pdf_path))

    document = {
        "document_id": document_id,
        "filename": filename,
        "pdf_path": str(pdf_path),
        "created_at": datetime.now().isoformat(),
        "status": "uploaded",
    }
    review_store.save_document(document)
    return jsonify({"success": True, "data": document})


@app.route('/api/sts/register-url', methods=['POST'])
def register_sts_pdf_url():
    data = request.get_json(silent=True) or {}
    pdf_url = (data.get("url") or "").strip()
    if not pdf_url:
        return jsonify({"success": False, "error": "Missing PDF URL"}), 400
    if not pdf_url.lower().startswith(("http://", "https://")):
        return jsonify({"success": False, "error": "PDF URL must start with http:// or https://"}), 400

    document_id = f"sts_{uuid.uuid4().hex[:12]}"
    filename = secure_filename(data.get("filename") or Path(pdf_url.split("?")[0]).name or f"{document_id}.pdf")
    if not filename.lower().endswith(".pdf"):
        filename = f"{filename}.pdf"
    document = {
        "document_id": document_id,
        "filename": filename,
        "pdf_url": pdf_url,
        "pdf_path": "",
        "created_at": datetime.now().isoformat(),
        "status": "url_registered",
    }
    review_store.save_document(document)
    return jsonify({"success": True, "data": document})


@app.route('/api/sts/parse/<document_id>', methods=['POST'])
def parse_sts_pdf(document_id):
    document = review_store.get_document(document_id)
    if not document:
        return jsonify({"success": False, "error": "Document not found"}), 404
    try:
        extracted = skill_registry.run("sts.extract_text_from_pdf", {
            "document_id": document_id,
            "pdf_path": document.get("pdf_path", ""),
            "pdf_url": document.get("pdf_url", ""),
        })
        split = skill_registry.run("sts.split_sections", {
            "document_id": document_id,
            "pages": extracted["pages"],
            "mineru_artifacts": extracted.get("mineru_artifacts", {}),
        })
        knowledge = skill_registry.run("sts.build_requirement_knowledge", {
            "document_id": document_id,
            "sections": split["sections"],
            "mineru_artifacts": extracted.get("mineru_artifacts", {}),
        })
        clusters = skill_registry.run("sts.cluster_topics", {
            "document_id": document_id,
            "sections": split["sections"],
            "requirement_units": knowledge["requirement_units"],
        })
        document.update({
            "status": "parsed",
            "pages": extracted["pages"],
            "sections": split["sections"],
            "requirement_units": knowledge["requirement_units"],
            "signals": knowledge["signals"],
            "coverage_matrix": knowledge["coverage_matrix"],
            "topic_clusters": clusters["topic_clusters"],
            "mineru_artifacts": extracted.get("mineru_artifacts", {}),
            "knowledge_artifacts": knowledge.get("knowledge_artifacts", {}),
            "cluster_artifacts": clusters.get("cluster_artifacts", {}),
            "parsed_at": datetime.now().isoformat(),
        })
        review_store.save_document(document)
        return jsonify({"success": True, "data": document})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/sts/generate-cases/<document_id>', methods=['POST'])
def generate_sts_cases(document_id):
    document = review_store.get_document(document_id)
    if not document:
        return jsonify({"success": False, "error": "Document not found"}), 404
    data = request.get_json(silent=True) or {}
    section_ids = set(data.get("section_ids") or [])
    cluster_ids = set(data.get("cluster_ids") or [])
    sections = document.get("sections") or []
    requirement_units = document.get("requirement_units") or []
    clusters = document.get("topic_clusters") or []
    if cluster_ids:
        clusters = [cluster for cluster in clusters if cluster.get("cluster_id") in cluster_ids]
        sections = []
        requirement_units = []
        for cluster in clusters:
            sections.extend(cluster.get("sections") or [])
            requirement_units.extend(cluster.get("requirement_units") or [])
    if section_ids:
        sections = [sec for sec in sections if sec.get("section_id") in section_ids]
        requirement_units = [unit for unit in requirement_units if unit.get("section_id") in section_ids]
        clusters = []
    if not sections and not clusters and not requirement_units:
        return jsonify({"success": False, "error": "No sections available for generation"}), 400
    try:
        requested_max_units = data.get("max_units")
        if requested_max_units in (None, "", 0) and cluster_ids:
            requested_max_units = os.getenv("STS_MAX_UNITS_PER_GENERATION", 3)
        generated = skill_registry.run("sts.generate_cases_from_sections", {
            "document_id": document_id,
            "sections": sections,
            "requirement_units": requirement_units,
            "topic_clusters": clusters,
            "signals": document.get("signals") or [],
            "model": data.get("model") or os.getenv("STS_GENERATION_MODEL", "qwen-plus"),
            "max_units": requested_max_units,
            "max_sections": data.get("max_sections"),
            "use_rag": data.get("use_rag", True),
            "testcase_count": data.get("testcase_count") or os.getenv("STS_CASES_PER_REQUIREMENT", 2),
        })
        pending = review_store.add_pending_cases(generated["generated_cases"])
        return jsonify({"success": True, "data": {"document_id": document_id, "generated_cases": pending}})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/sts/artifacts/<document_id>', methods=['GET'])
def get_sts_artifacts(document_id):
    document = review_store.get_document(document_id)
    if not document:
        return jsonify({"success": False, "error": "Document not found"}), 404
    return jsonify({
        "success": True,
        "data": {
            "document_id": document_id,
            "filename": document.get("filename", ""),
            "mineru_artifacts": document.get("mineru_artifacts", {}),
            "knowledge_artifacts": document.get("knowledge_artifacts", {}),
            "cluster_artifacts": document.get("cluster_artifacts", {}),
            "counts": {
                "sections": len(document.get("sections") or []),
                "requirement_units": len(document.get("requirement_units") or []),
                "signals": len(document.get("signals") or []),
                "topic_clusters": len(document.get("topic_clusters") or []),
            },
        },
    })


@app.route('/api/review/test-cases', methods=['GET'])
def list_review_cases():
    status = request.args.get("status", "pending")
    return jsonify({"success": True, "data": review_store.list_cases(status)})


@app.route('/api/review/test-cases/<case_id>', methods=['POST'])
def review_test_case(case_id):
    data = request.get_json(silent=True) or {}
    try:
        result = review_store.review_case(
            case_id=case_id,
            decision=data.get("decision", ""),
            reviewer=data.get("reviewer", ""),
            comment=data.get("comment", ""),
        )
        return jsonify({"success": True, "data": result})
    except KeyError as e:
        return jsonify({"success": False, "error": str(e)}), 404
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400


@app.route('/api/review/generated-case', methods=['POST'])
def save_generated_case_review():
    data = request.get_json(silent=True) or {}
    decision = data.get("decision", "")
    try:
        result = review_store.save_reviewed_case(
            case=data.get("case", {}),
            decision=decision,
            reviewer=data.get("reviewer", ""),
            comment=data.get("comment", ""),
        )
        return jsonify({"success": True, "data": result})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400


@app.route('/api/test-cases/by-type-expected', methods=['GET'])
def get_cases_by_type_expected():
    """Get cases grouped by test type and expected behavior (two-level)."""
    try:
        # Load cases
        case_classifier.load_cases()
        
        # Group by test type and expected behavior
        type_expected_groups = {}
        
        for case in case_classifier.test_cases:
            original_data = case.get('original_data', {})
            test_type = original_data.get('测试用例类型', '未分类')
            expected = original_data.get('预期行为', 'Unknown expected behavior')
            
            # Init test type bucket
            if test_type not in type_expected_groups:
                type_expected_groups[test_type] = {
                    'expected_behaviors': {}
                }
            
            # Init expected-behavior bucket
            if expected not in type_expected_groups[test_type]['expected_behaviors']:
                type_expected_groups[test_type]['expected_behaviors'][expected] = []
            
            # Build case record
            case_info = {
                'unique_id': case.get('unique_id', f"ROW_{case.get('row_number', 0)}"),
                '编号': original_data.get('编号', ''),
                '测试用例名称': original_data.get('测试用例名称', ''),
                '测试用例类型': test_type,
                '前提条件': case.get('前提条件_分割', []),
                '执行动作': original_data.get('执行动作', ''),
                '预期行为': expected,
                'row_number': case.get('row_number', 0),
                '是否单一预期': case_classifier._is_single_expected(expected)
            }
            
            type_expected_groups[test_type]['expected_behaviors'][expected].append(case_info)
        
        # Statistics
        total_cases = len(case_classifier.test_cases)
        total_types = len(type_expected_groups)
        total_expected = sum(len(data['expected_behaviors']) for data in type_expected_groups.values())
        
        statistics = {
            "总案例数": total_cases,
            "测试类型数": total_types,
            "预期行为数": total_expected,
            "平均每类型预期数": round(total_expected / total_types, 1) if total_types else 0
        }
        
        print(f"Type-expected classification done: {total_types} types, {total_expected} expected behaviors, {total_cases} cases")
        
        return jsonify({
            "success": True,
            "data": type_expected_groups,
            "statistics": statistics
        })
    except Exception as e:
        print(f"Failed to get type-expected classification: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


def _evaluate_testcase_internal(text):
    try:
        result = skill_registry.run("testcase.evaluate", {"testcase": text})
        return result["score"], result["issues"]
    except Exception as skill_err:
        print(f"Skill testcase.evaluate failed, using legacy evaluator: {type(skill_err).__name__}: {skill_err}")

    """Return (score, issues). Shared by API route and auto-improve loop."""
    try:
        prompt = f"""You are an automotive test case evaluation expert. Assess whether the generated test case below is reasonable, executable, and verifiable; give improvement suggestions.

Requirements:
1) Output a single JSON object with fields:
   - is_reasonable: boolean
   - score: integer 0-100
   - confidence: float 0-1
   - issues: string array
   - suggestions: string array
2) Keep each issue/suggestion concise for UI display.
3) If structure is missing (测试项/前提条件/执行动作/预期行为), state it clearly.

【Generated test case】
{text}
"""
        critic_resp, _ = generate_testcase(prompt, model="Qwen3-1.7B-Critic")
        candidate = (critic_resp or "").strip()
        m = re.search(r"\{[\s\S]*\}", candidate)
        if m:
            candidate = m.group(0)
        result = json.loads(candidate)
        is_reasonable = bool(result.get("is_reasonable"))
        score = max(0, min(100, int(result.get("score", 0))))
        confidence = max(0.0, min(1.0, float(result.get("confidence", 0.0))))
        issues = [str(x) for x in (result.get("issues") or []) if str(x).strip()]
        suggestions = [str(x) for x in (result.get("suggestions") or []) if str(x).strip()]
        merged = issues + ([f"Overall: {'reasonable' if is_reasonable else 'not reasonable'} (confidence {confidence:.2f})"] if confidence > 0 else [])
        merged += [f"Suggestion: {s}" for s in suggestions]
        return score, merged
    except Exception:
        pass
    issues = []
    score = 100
    for sec in ["测试项", "前提条件", "执行动作", "预期行为"]:
        if sec not in text:
            issues.append(f"Missing field: {sec}")
            score -= 20
    m = re.search(r"前提条件\s*[:：]\s*(.+?)(\n\s*执行动作\s*[:：]|\n\s*预期行为\s*[:：]|$)", text, flags=re.S)
    if m:
        if not re.search(r"(^|\n)\s*\d+\s*[\.|、]\s*\S+", m.group(1).strip()):
            issues.append("Preconditions should use a numbered list (e.g. 1. / 1、)")
            score -= 10
    elif "前提条件" in text:
        issues.append("Could not parse preconditions")
        score -= 10
    if not bool(re.search(r"执行动作\s*[:：]\s*\S+", text)) and "执行动作" in text:
        issues.append("Actions empty or too short")
        score -= 10
    if not bool(re.search(r"预期行为\s*[:：]\s*\S+", text)) and "预期行为" in text:
        issues.append("Expected behavior empty or too short")
        score -= 10
    if len(text) < 120:
        issues.append("Content too short; add actionable steps and verifiable expectations")
        score -= 10
    return max(0, min(100, score)), issues


@app.route('/api/evaluate-testcase', methods=['POST'])
def evaluate_testcase():
    try:
        data = request.get_json(silent=True) or {}
        text = (data.get('testcase') or '').strip()
        if not text:
            return jsonify({"success": False, "error": "Missing testcase content"}), 400
        score, issues = _evaluate_testcase_internal(text)
        return jsonify({"success": True, "score": score, "issues": issues})
    except Exception as e:
        print(f"Evaluation failed: {str(e)}")
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/generate-reverse-from-expected', methods=['POST'])
def generate_reverse_from_expected():
    """Generate reverse test cases from expected behavior."""
    try:
        data = request.get_json()
        cases = data.get('cases', [])
        
        if not cases:
            return jsonify({"success": False, "error": "No cases provided"}), 400
        
        print("\nStarting reverse case generation from expected behavior")
        print(f"Received {len(cases)} case(s)")
        
        # Step 1: filter single positive expected-behavior cases
        single_expected_cases = []
        for case in cases:
            # Must be single expected behavior
            if not case.get('是否单一预期', False):
                continue
            
            # Expected behavior must be non-empty
            expected = case.get('预期行为', '').strip()
            if not expected:
                continue
            
            # Exclude clearly negative expectations
            negative_keywords = ['不', '无', '失败', '异常', '错误', '禁止']
            is_negative = any(keyword in expected for keyword in negative_keywords)
            
            # Keep only positive-expectation cases
            if not is_negative:
                single_expected_cases.append(case)
                print(f"   Kept positive case: {case.get('测试用例名称', '')[:40]}... | expected: {expected[:30]}...")
            else:
                print(f"   Skipped negative case: {case.get('测试用例名称', '')[:40]}... | expected: {expected[:30]}...")
        
        print(f"\nFiltered to {len(single_expected_cases)} single positive-expectation case(s)")
        
        if len(single_expected_cases) == 0:
            return jsonify({
                "success": False,
                "error": "No matching single positive-expectation cases"
            }), 400
        
        # Step 2: group by expected behavior and action
        groups = {}
        for case in single_expected_cases:
            expected = case.get('预期行为', '')
            action = case.get('执行动作', '')
            key = f"{expected}|||{action}"  # Composite key separator
            
            if key not in groups:
                groups[key] = []
            groups[key].append(case)
        
        print(f"Grouped into {len(groups)} bucket(s) (same expected + action)")
        
        # Step 3: key preconditions → reverse cases
        all_generated_cases = []
        
        for key, group_cases in groups.items():
            expected, action = key.split('|||')
            print(f"\nProcessing group: {len(group_cases)} case(s)")
            print(f"   Expected: {expected[:50]}...")
            print(f"   Action: {action[:50]}...")
            
            # Extract key preconditions shared by all cases in the group
            if len(group_cases) == 1:
                # Single case: all preconditions are key
                key_preconditions = group_cases[0].get('前提条件', [])
            else:
                # Multiple cases: intersection
                precondition_sets = [set(case.get('前提条件', [])) for case in group_cases]
                key_preconditions = list(set.intersection(*precondition_sets))
            
            print(f"   Key preconditions: {len(key_preconditions)}")
            
            if len(key_preconditions) == 0:
                print("   No key preconditions; skipping")
                continue
            
            # For each key precondition, generate reverse case
            for precondition in key_preconditions:
                # Negate via LLM
                reversed_precondition = reverse_precondition_with_llm(precondition)
                reversed_expected = reverse_expected_with_llm(expected)
                
                # Build reverse case
                reverse_case = {
                    '测试用例名称': f"Reverse-{group_cases[0].get('测试用例名称', '')}",
                    '测试用例类型': group_cases[0].get('测试用例类型', ''),
                    '前提条件': [reversed_precondition if p == precondition else p for p in group_cases[0].get('前提条件', [])],
                    '执行动作': action,
                    '预期行为': reversed_expected,
                    '生成方式': 'Key precondition negation from expected behavior',
                    '原始前提条件': precondition,
                    '反向前提条件': reversed_precondition
                }
                
                all_generated_cases.append(reverse_case)
        
        print(f"\nGenerated {len(all_generated_cases)} reverse case(s)")
        
        return jsonify({
            "success": True,
            "generated_count": len(all_generated_cases),
            "generated_cases": all_generated_cases,
            "message": f"Generated {len(all_generated_cases)} reverse test case(s)"
        })
        
    except Exception as e:
        print(f"Reverse case generation failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

def reverse_precondition_with_llm(precondition: str) -> str:
    """Negate a precondition using the LLM."""
    prompt = f"""You are a test case expert. Logically negate the precondition below.

Original precondition: {precondition}

Negation rules (Chinese domain examples follow):
1. State: on→off, valid→invalid, normal→abnormal, connected→disconnected
2. Time: day↔night, workday↔rest day
3. Environment: indoor↔outdoor, sunny↔rainy, hot↔cold
4. Numeric: greater→less-or-equal, satisfied→not satisfied, reached→not reached
5. Existence: exists→not exists, has→none, done→not done

Principles:
- Understand semantics; pick natural opposites
- For time/environment use natural pairs (e.g. 夜晚 → 白天)
- Do not only prefix 不; use idiomatic Chinese
- Keep output concise and natural in Chinese

Example 1:
Original: 设置伴我回家功能开启
Negated: 设置伴我回家功能关闭

Example 2:
Original: 夜晚
Negated: 白天
Note: opposite of 夜晚 is 白天

Example 3:
Original: 车辆处于P档
Negated: 车辆不处于P档

Example 4:
Original: 车速大于5km/h
Negated: 车速小于等于5km/h

Example 5:
Original: 电池电量充足
Negated: 电池电量不足

Example 6:
Original: 门已关闭
Negated: 门未关闭

Negate the precondition above; output one natural Chinese result only:"""

    try:
        result = generate_testcase(prompt, model="qwen3-8b")
        
        if isinstance(result, tuple):
            response, elapsed_time = result
        else:
            response = result
            elapsed_time = 0
        
        # Keep first line of LLM output
        reversed_condition = response.strip().split('\n')[0].strip()
        print(f"      Negated precondition: {precondition} → {reversed_condition} ({elapsed_time:.2f}s)")
        return reversed_condition
    except Exception as e:
        import traceback
        print(f"      LLM negation failed: {e}")
        traceback.print_exc()
        return f"NOT({precondition})"

def reverse_expected_with_llm(expected: str) -> str:
    """Negate expected behavior using the LLM."""
    prompt = f"""You are a test case expert. Logically negate the expected behavior below.

Original expected behavior: {expected}

Rules:
1. Negate only the primary behavior, not follow-on timing/conditions
2. If multiple actions, negate the first core action only
3. Execute → does not execute; works → does not work; on/display → off/not displayed
4. Success → failure; valid → invalid; trigger/activate → no trigger/activate

Principles:
- When the main function fails, downstream effects need not be negated separately
- Output core negation only; no extra duration or state
- Keep Chinese concise; do not add details not in the original

Example 1:
Original: 伴我回家功能激活
Negated: 伴我回家功能不激活

Example 2:
Original: 伴我回家功能激活，5s、10s、20s后熄灭
Negated: 伴我回家功能不激活

Example 3:
Original: 尾门打开到记忆高度
Negated: 尾门不打开

Example 4:
Original: 灯光点亮并保持30秒
Negated: 灯光不点亮

Example 5:
Original: 会触发欢送灯光
Negated: 不触发欢送灯光

Negate the expected behavior above; output core Chinese negation only:"""

    try:
        result = generate_testcase(prompt, model="qwen3-8b")
        
        if isinstance(result, tuple):
            response, elapsed_time = result
        else:
            response = result
            elapsed_time = 0
        
        # Keep first line of LLM output
        reversed_expected = response.strip().split('\n')[0].strip()
        print(f"      Negated expected: {expected[:30]}... → {reversed_expected[:30]}... ({elapsed_time:.2f}s)")
        return reversed_expected
    except Exception as e:
        import traceback
        print(f"      LLM expected negation failed: {e}")
        traceback.print_exc()
        return f"NOT({expected})"

@app.route('/api/export-generated-cases-excel', methods=['POST'])
def export_generated_cases_excel():
    """Export generated reverse cases to Excel (case_generation format)."""
    try:
        data = request.get_json()
        cases = data.get('cases', [])
        
        if not cases:
            return jsonify({"success": False, "error": "No cases to export"}), 400
        
        # Excel rows (case_generation format)
        excel_data = []
        
        for case in cases:
            # Preconditions list → string
            preconditions = case.get('前提条件', [])
            if isinstance(preconditions, list):
                preconditions_str = '\n'.join([f"{i+1}. {cond}" for i, cond in enumerate(preconditions)])
            else:
                preconditions_str = str(preconditions)
            
            # Case name
            case_name = case.get('测试用例名称', '')
            
            # Expected behavior
            expected = case.get('预期行为', '')
            
            # Generation basis
            generation_basis = ''
            if case.get('原始前提条件') and case.get('反向前提条件'):
                generation_basis = f"Original: {case.get('原始前提条件')}\nNegated: {case.get('反向前提条件')}"
            elif case.get('生成方式'):
                generation_basis = case.get('生成方式')
            
            excel_data.append({
                '测试案例名称': case_name,
                '前提条件': preconditions_str,
                '执行动作': case.get('执行动作', ''),
                '预期行为': expected,
                '生成依据': generation_basis
            })
        
        # Build DataFrame
        df = pd.DataFrame(excel_data)
        
        # Write Excel
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Generated test cases')
            
            # Column widths (case_generation)
            worksheet = writer.sheets['Generated test cases']
            worksheet.column_dimensions['A'].width = 40  # case name
            worksheet.column_dimensions['B'].width = 50  # preconditions
            worksheet.column_dimensions['C'].width = 40  # actions
            worksheet.column_dimensions['D'].width = 40  # expected
            worksheet.column_dimensions['E'].width = 50  # basis
            
            # Wrap text
            from openpyxl.styles import Alignment
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        output.seek(0)
        
        # Output filename
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'generated_test_cases_{timestamp}.xlsx'
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        print(f"Excel export failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/test-cases/by-action', methods=['GET'])
def get_cases_by_action():
    """Get cases grouped by test type and action (two-level)."""
    try:
        # Load and classify cases
        case_classifier.load_cases()
        type_action_groups = case_classifier.classify_by_type_and_action()
        
        # Statistics
        total_cases = sum(group['total_count'] for group in type_action_groups.values())
        total_types = len(type_action_groups)
        total_actions = sum(len(group['actions']) for group in type_action_groups.values())
        
        statistics = {
            "总案例数": total_cases,
            "测试类型数": total_types,
            "动作分类数": total_actions,
            "平均每类型案例数": round(total_cases / total_types, 1) if total_types else 0
        }
        
        return jsonify({
            "success": True,
            "data": type_action_groups,
            "statistics": statistics
        })
    except Exception as e:
        print(f"Failed to get classified cases: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/test-cases/generate-reverse', methods=['POST'])
def generate_reverse_cases():
    """Generate reverse test cases."""
    try:
        data = request.get_json()
        case_ids = data.get('case_ids', [])
        generation_mode = data.get('generation_mode', 'single')
        model_name = data.get('model_name', 'qwen2.5-3b-finetune-data-continued')
        
        if not case_ids:
            return jsonify({"success": False, "error": "No cases selected"}), 400
        
        print("\nStarting test case generation")
        print(f"Generation mode: {generation_mode}")
        print(f"Selected cases: {len(case_ids)}")
        print(f"Model: {model_name}")
        
        # All cases
        case_classifier.load_cases()
        action_groups = case_classifier.classify_by_action()
        
        # Flatten cases
        all_cases = []
        for cases in action_groups.values():
            all_cases.extend(cases)
        
        # Create generation task
        task_id = reverse_generator.create_task(case_ids, all_cases, generation_mode)
        
        mode_text = "1:1 reverse cases" if generation_mode == 'single' else "batch composite cases"
        
        return jsonify({
            "success": True,
            "task_id": task_id,
            "message": f"Starting {mode_text}; please wait..."
        })
        
    except Exception as e:
        print(f"Case generation failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@app.route('/api/test-cases/retry/<task_id>', methods=['POST'])
def retry_reverse_task(task_id):
    """Retry a failed reverse-generation task when possible."""
    task = reverse_generator.get_task_status(task_id)
    if not task:
        return jsonify({"success": False, "error": "Task not found", "retryable": False}), 404
    cases = task.get('cases_snapshot') or []
    if not cases:
        return jsonify({
            "success": False,
            "error": "No retry snapshot; re-select cases and generate again",
            "retryable": False,
        }), 400
    case_ids = [case.get('unique_id') or case.get('编号') for case in cases if case.get('unique_id') or case.get('编号')]
    new_task_id = reverse_generator.create_task(case_ids, cases, task.get('mode') or 'batch')
    return jsonify({"success": True, "task_id": new_task_id, "message": "Generation task restarted"})


@app.route('/api/test-cases/generation-status/<task_id>', methods=['GET'])
def get_generation_status(task_id):
    """Query reverse case generation progress."""
    try:
        task = reverse_generator.get_task_status(task_id)
        
        if not task:
            return jsonify({"success": False, "error": "Task not found"}), 404
        
        return jsonify({
            "success": True,
            "status": task['status'],
            "progress": {
                "total": task['total'],
                "completed": task['completed'],
                "failed": task['failed'],
                "percentage": task.get('percentage', 0)
            },
            "results": task['results'][:5]
        })
        
    except Exception as e:
        print(f"Progress query failed: {str(e)}")
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/test-cases/reverse-results/<task_id>', methods=['GET'])
def get_reverse_results(task_id):
    """Get full reverse case generation results."""
    try:
        results = reverse_generator.get_task_results(task_id)
        
        if not results:
            return jsonify({"success": False, "error": "Task not found"}), 404
        
        return jsonify({
            "success": True,
            "data": results
        })
        
    except Exception as e:
        print(f"Failed to get results: {str(e)}")
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/test-cases/export-excel/<task_id>', methods=['GET'])
def export_to_excel(task_id):
    """Export generated cases to Excel."""
    try:
        results = reverse_generator.get_task_results(task_id)
        
        if not results:
            return jsonify({"success": False, "error": "Task not found"}), 404
        
        # Original cases
        case_classifier.load_cases()
        action_groups = case_classifier.classify_by_action()
        all_cases = []
        for cases in action_groups.values():
            all_cases.extend(cases)
        
        # Map case id → name
        case_id_to_name = {case['编号']: case['测试用例名称'] for case in all_cases}
        
        # Excel rows
        excel_data = []
        
        for case in results.get('reverse_cases', []):
            related_ids = case.get('关联原始案例', [])
            if isinstance(related_ids, list) and len(related_ids) > 0:
                related_names = [case_id_to_name.get(rid, rid) for rid in related_ids]
                generation_basis = '、'.join(related_names)
            else:
                original_id = case.get('原始案例编号', '')
                generation_basis = case_id_to_name.get(original_id, original_id) if original_id else ''
            
            preconditions = case.get('前提条件', case.get('反向前提条件', []))
            if isinstance(preconditions, list):
                preconditions_str = '\n'.join([f"{i+1}. {cond}" for i, cond in enumerate(preconditions)])
            else:
                preconditions_str = str(preconditions)
            
            expected = case.get('预期行为', case.get('反向预期行为', ''))
            case_name = case.get('案例名称', case.get('反向案例名称', ''))
            
            excel_data.append({
                '测试案例名称': case_name,
                '前提条件': preconditions_str,
                '执行动作': case.get('执行动作', ''),
                '预期行为': expected,
                '生成依据': generation_basis
            })
        
        # Build DataFrame
        df = pd.DataFrame(excel_data)
        
        # Write Excel
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Generated test cases')
            
            worksheet = writer.sheets['Generated test cases']
            worksheet.column_dimensions['A'].width = 40
            worksheet.column_dimensions['B'].width = 50
            worksheet.column_dimensions['C'].width = 40
            worksheet.column_dimensions['D'].width = 40
            worksheet.column_dimensions['E'].width = 50
            
            from openpyxl.styles import Alignment
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        output.seek(0)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'generated_test_cases_{timestamp}.xlsx'
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        print(f"Excel export failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/generate', methods=['POST'])
def generate():
    """Generate test cases from the main UI."""
    try:
        user_input = request.form.get('user_input', '').strip()
        testcase_count = request.form.get('testcase_count', '1').strip()
        
        if not user_input:
            return jsonify({'error': 'Please enter a valid test requirement.'})
        
        try:
            testcase_count = int(testcase_count)
            if testcase_count < 1 or testcase_count > 10:
                return jsonify({'error': 'Test case count must be between 1 and 10.'})
        except ValueError:
            return jsonify({'error': 'Test case count must be a valid integer.'})

        print(f"\nUser request: {user_input}")
        print(f"Test case count: {testcase_count}")
        
        skill_result = skill_registry.run("testcase.generate_from_requirement", {
            "user_input": user_input,
            "testcase_count": testcase_count,
            "model": "Qwen3-4B",
        })
        testcase = skill_result["testcase"]
        triples = skill_result["triples"]
        similar_cases = skill_result["similar_cases"]
        pdf_kg_evidence = skill_result.get("pdf_kg_evidence", [])
        elapsed_time = skill_result["elapsed_time"]
        
        print("Test case generation succeeded")
        print(f"Total time: {elapsed_time:.2f}s")

        formatted_triples = []
        nodes = set()
        edges = []
        for triple in triples:
            processed = process_triple(triple.copy())
            head = processed['head']
            tail = processed['tail']
            relation = processed['relation']
            nodes.add(head)
            nodes.add(tail)
            edges.append({'from': head, 'to': tail, 'label': relation})

            formatted_triples.append({
                'head': wrap_text(head, 25),
                'relation': wrap_text(relation, 15),
                'tail': wrap_text(tail, 25)
            })

        node_list = [{'id': node, 'label': node} for node in nodes]

        return jsonify({
            'testcase': testcase,
            'triples': formatted_triples,
            'similar_cases': similar_cases,
            'pdf_kg_evidence': pdf_kg_evidence,
            'graph': {
                'nodes': node_list,
                'edges': edges
            }
        })
    except Exception as e:
        print(f"\nError: {type(e).__name__}")
        print(f"Error message: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Generation failed: {str(e)}'}), 500


@app.route('/generate-stream', methods=['POST'])
def generate_stream():
    """Generate test cases from the main UI (stream thoughts + final result)."""
    data = request.get_json(silent=True) or {}
    user_input = (data.get('user_input') or '').strip()

    if not user_input:
        return jsonify({'error': 'Please enter a valid test requirement.'}), 400

    testcase_count = 1
    event_queue = queue.Queue()

    def worker():
        try:
            print(f"\nUser request: {user_input}")
            print(f"Test case count: {testcase_count}")

            def publish_progress(event):
                if isinstance(event, dict):
                    event_queue.put(event)
                else:
                    event_queue.put({"type": "thought", "message": str(event)})

            skill_result = skill_registry.run("testcase.generate_from_requirement", {
                "user_input": user_input,
                "testcase_count": testcase_count,
                "model": "Qwen3-4B",
                "progress_callback": publish_progress,
            })
            testcase = skill_result["testcase"]
            similar_cases = skill_result["similar_cases"]
            pdf_kg_evidence = skill_result.get("pdf_kg_evidence", [])
            elapsed_time = skill_result["elapsed_time"]
            complete_testcases = skill_result["complete_testcases"]

            # Color per retrieved testcase
            CASE_COLORS = [
                {'background': '#dbeafe', 'border': '#3b82f6', 'highlight': '#93c5fd'},
                {'background': '#dcfce7', 'border': '#22c55e', 'highlight': '#86efac'},
                {'background': '#fef3c7', 'border': '#f59e0b', 'highlight': '#fcd34d'},
                {'background': '#fce7f3', 'border': '#ec4899', 'highlight': '#f9a8d4'},
                {'background': '#e0e7ff', 'border': '#6366f1', 'highlight': '#a5b4fc'},
                {'background': '#ccfbf1', 'border': '#14b8a6', 'highlight': '#5eead4'},
                {'background': '#fee2e2', 'border': '#ef4444', 'highlight': '#fca5a5'},
                {'background': '#f3e8ff', 'border': '#a855f7', 'highlight': '#d8b4fe'},
            ]

            # Dedupe list preserving order
            def _uniq_list(items):
                seen = set()
                out = []
                for it in items or []:
                    if not it:
                        continue
                    if it in seen:
                        continue
                    seen.add(it)
                    out.append(it)
                return out

            # Filter and dedupe testcases:
            # - Drop cases missing actions or expected behavior
            # - Dedupe by (name, actions, expectations)
            filtered_testcases = []
            seen_keys = set()
            for tc in (complete_testcases or []):
                name = (tc.get('case_name') or '').strip()
                pre = _uniq_list(tc.get('preconditions', []))
                acts = _uniq_list(tc.get('actions', []))
                exps = _uniq_list(tc.get('expected_behaviors', []))

                # Skip if actions or expectations empty
                if not acts or not exps:
                    continue

                key = (name.lower(), tuple(sorted(pre)), tuple(sorted(acts)), tuple(sorted(exps)))
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                filtered_testcases.append({
                    'case_name': name,
                    'test_type': tc.get('test_type'),
                    'preconditions': pre,
                    'actions': acts,
                    'expected_behaviors': exps,
                })

            # Testcase-centric graph
            node_set = {}
            edge_set = set()
            edge_list = []
            retrieved_testcases_data = []

            for idx, tc in enumerate(filtered_testcases):
                color = CASE_COLORS[idx % len(CASE_COLORS)]
                tc_name = tc.get('case_name', f'TestCase_{idx+1}')

                # TestCase hub node
                node_set[tc_name] = {
                    'id': tc_name,
                    'label': tc_name,
                    'color': color,
                    'group': idx,
                    'isTestCase': True,
                }

                # Precondition nodes/edges
                for item in tc.get('preconditions', []):
                    if item not in node_set:
                        node_set[item] = {'id': item, 'label': item, 'color': color, 'group': idx}
                    e = (tc_name, item, '前提条件')
                    if e not in edge_set:
                        edge_set.add(e)
                        edge_list.append({'from': tc_name, 'to': item, 'label': '前提条件', 'group': idx})

                # Action nodes/edges
                for item in tc.get('actions', []):
                    if item not in node_set:
                        node_set[item] = {'id': item, 'label': item, 'color': color, 'group': idx}
                    e = (tc_name, item, '执行动作')
                    if e not in edge_set:
                        edge_set.add(e)
                        edge_list.append({'from': tc_name, 'to': item, 'label': '执行动作', 'group': idx})

                # Expected-behavior nodes/edges
                for item in tc.get('expected_behaviors', []):
                    if item not in node_set:
                        node_set[item] = {'id': item, 'label': item, 'color': color, 'group': idx}
                    e = (tc_name, item, '预期行为')
                    if e not in edge_set:
                        edge_set.add(e)
                        edge_list.append({'from': tc_name, 'to': item, 'label': '预期行为', 'group': idx})

                # Card payload for frontend
                retrieved_testcases_data.append({
                    'case_name': tc_name,
                    'test_type': tc.get('test_type'),
                    'preconditions': tc.get('preconditions', []),
                    'actions': tc.get('actions', []),
                    'expected_behaviors': tc.get('expected_behaviors', []),
                    'color': color['border'],
                })

            event_queue.put(
                {
                    "type": "result",
                    "data": {
                        'testcase': testcase,
                        'retrieved_testcases': retrieved_testcases_data,
                        'similar_cases': similar_cases,
                        'pdf_kg_evidence': pdf_kg_evidence,
                        'graph': {
                            'nodes': list(node_set.values()),
                            'edges': edge_list,
                        },
                        'elapsed_time': elapsed_time,
                    },
                }
            )
        except Exception as e:
            print(f"\nError: {type(e).__name__}")
            print(f"Error message: {str(e)}")
            import traceback
            traceback.print_exc()
            event_queue.put({"type": "error", "message": f"Generation failed: {str(e)}"})
        finally:
            event_queue.put({"type": "done"})

    threading.Thread(target=worker, daemon=True).start()

    @stream_with_context
    def generate_events():
        while True:
            event = event_queue.get()
            yield json.dumps(event, ensure_ascii=False) + "\n"
            if event.get("type") == "done":
                break

    return Response(generate_events(), mimetype='application/x-ndjson')


@app.route('/auto-improve-stream', methods=['POST'])
def auto_improve_stream():
    """Hermes Agent self-improvement loop: evaluate → refine until threshold or max iterations."""
    data = request.get_json(silent=True) or {}
    current_testcase = (data.get('current_testcase') or '').strip()
    threshold = max(50, min(95, int(data.get('threshold', 80))))
    max_iterations = max(1, min(5, int(data.get('max_iterations', 3))))

    if not current_testcase:
        return jsonify({'error': 'Missing current testcase content'}), 400

    event_queue = queue.Queue()

    def worker():
        nonlocal current_testcase
        best_score = 0
        best_testcase = current_testcase
        try:
            event_queue.put({"type": "auto_start", "threshold": threshold, "max_iterations": max_iterations})

            for i in range(max_iterations):
                iteration = i + 1

                # Think: plan iteration
                event_queue.put({"type": "auto_phase", "phase": "evaluating",
                                  "iteration": iteration, "message": f"Round {iteration}: evaluating..."})

                # Act: evaluate
                score, issues = _evaluate_testcase_internal(current_testcase)

                # Observe: record result
                event_queue.put({"type": "auto_eval", "iteration": iteration,
                                  "score": score, "issues": issues,
                                  "reached": score >= threshold})

                if score > best_score:
                    best_score = score
                    best_testcase = current_testcase

                if score >= threshold:
                    event_queue.put({"type": "auto_phase", "phase": "done",
                                      "iteration": iteration,
                                      "message": f"Round {iteration} met threshold ({score} ≥ {threshold}); stopping"})
                    break

                if i < max_iterations - 1:
                    # Think: plan refinement
                    issues_text = '\n'.join(f'  - {x}' for x in issues)
                    event_queue.put({"type": "auto_phase", "phase": "refining",
                                      "iteration": iteration,
                                      "message": f"Round {iteration}: refining ({score} < {threshold} target)..."})

                    refine_prompt = f"""You are an automotive test expert. Improve the testcase from evaluation feedback; output the full improved case only.

【Current testcase ({score} points)】
{current_testcase}

【Issues (target {threshold} points)】
{issues_text}

Keep the test goal; retain sections 测试项/前提条件/执行动作/预期行为; no explanation."""


                    # Act: refine
                    refine_result = skill_registry.run("testcase.refine", {
                        "current_testcase": current_testcase,
                        "instruction": refine_prompt,
                        "eval_issues": issues,
                        "model": "Qwen3-4B",
                    })
                    refined_raw = refine_result.get("testcase", "")
                    refined_clean = re.sub(r'<think>[\s\S]*?</think>', '', refined_raw or '', flags=re.IGNORECASE).strip()
                    if refined_clean:
                        current_testcase = refined_clean
                else:
                    event_queue.put({"type": "auto_phase", "phase": "max_reached",
                                      "iteration": iteration,
                                      "message": f"Max iterations reached; best score: {best_score}"})

            event_queue.put({
                "type": "result",
                "data": {
                    'testcase': best_testcase,
                    'final_score': best_score,
                    'refine_mode': True,
                }
            })
        except Exception as e:
            print(f"\nAuto-improve failed: {type(e).__name__}: {str(e)}")
            import traceback
            traceback.print_exc()
            event_queue.put({"type": "error", "message": f"Auto-improve failed: {str(e)}"})
        finally:
            event_queue.put({"type": "done"})

    threading.Thread(target=worker, daemon=True).start()

    @stream_with_context
    def generate_events():
        while True:
            event = event_queue.get()
            yield json.dumps(event, ensure_ascii=False) + "\n"
            if event.get("type") == "done":
                break

    return Response(generate_events(), mimetype='application/x-ndjson')


@app.route('/refine-stream', methods=['POST'])
def refine_stream():
    """Refine the current testcase directly (skip KG retrieval; use eval feedback + instructions)."""
    data = request.get_json(silent=True) or {}
    current_testcase = (data.get('current_testcase') or '').strip()
    instruction = (data.get('instruction') or '').strip()
    eval_issues = data.get('eval_issues') or []

    if not current_testcase:
        return jsonify({'error': 'Missing current testcase content'}), 400

    event_queue = queue.Queue()

    def worker():
        try:
            issues_text = ''
            if eval_issues:
                issues_text = '\n\nEvaluation issues and suggestions:\n' + '\n'.join(f'  - {x}' for x in eval_issues)

            instruction_text = instruction or 'Apply evaluation feedback and produce an improved version'

            prompt = f"""You are an automotive test expert. Improve the testcase below; output the full improved case only.

【Current testcase】
{current_testcase}{issues_text}

【Requirements】
{instruction_text}

Notes:
- Keep the original test goal
- Keep section format 测试项/前提条件/执行动作/预期行为
- Output the improved case only; no explanation"""


            event_queue.put({"type": "thought", "message": "Direct refine mode (skipping KG retrieval)"})
            if eval_issues:
                event_queue.put({"type": "thought", "message": f"Using {len(eval_issues)} evaluation note(s)"})
            event_queue.put({"type": "thought", "message": f"Refine instruction: {instruction_text}"})

            refine_result = skill_registry.run("testcase.refine", {
                "current_testcase": current_testcase,
                "instruction": instruction_text,
                "eval_issues": eval_issues,
                "model": "Qwen3-4B",
            })
            testcase = refine_result.get("testcase", "")
            elapsed_time = refine_result.get("elapsed_time", 0)

            event_queue.put({
                "type": "result",
                "data": {
                    'testcase': testcase,
                    'elapsed_time': elapsed_time,
                    'refine_mode': True,
                }
            })
        except Exception as e:
            print(f"\nRefine failed: {type(e).__name__}: {str(e)}")
            import traceback
            traceback.print_exc()
            event_queue.put({"type": "error", "message": f"Refine failed: {str(e)}"})
        finally:
            event_queue.put({"type": "done"})

    threading.Thread(target=worker, daemon=True).start()

    @stream_with_context
    def generate_events():
        while True:
            event = event_queue.get()
            yield json.dumps(event, ensure_ascii=False) + "\n"
            if event.get("type") == "done":
                break

    return Response(generate_events(), mimetype='application/x-ndjson')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=4000, debug=True)
